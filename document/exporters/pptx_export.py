"""شرائح من الخطة نفسها — عرضٌ للمعلّم لا نسخةٌ ثانية من المستند.

المستند للقراءة، والشرائح للعرض؛ وما يصلح لأحدهما يقتل الآخر. لذلك لا
تُنسخ الفقرات هنا: كل قسمٍ يصير شريحةً واحدة بعنوانه ونقاطه المكثّفة،
وكل لقطةٍ شاشة تصير شريحةً كاملة بتسميتها. النصّ الكامل يبقى حيث مكانه
الطبيعي — في ملاحظات المتحدّث أسفل الشريحة، يقرؤها هو ولا تُعرض.

**‏python-pptx اعتماد اختياري.** غيابها يتخطّى هذه الصيغة وحدها بسطر
في السجلّ ولا يمسّ شيئًا آخر — نفس عقد بقية المُصيِّرات.

**اتجاه النصّ يُضبط بـ XML مباشرةً.** ‏python-pptx لا تعرض خاصيّة
لاتجاه الفقرة إطلاقًا، والشريحة العربية بلا ``rtl="1"`` تخرج بعلامات
الترقيم في أوّل السطر بدل آخره — عطلٌ يظهر للجمهور لا للمطوّر.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional

from config.schemas import DocumentSection
from document.alt_text import describe
from document.exporters import ExportContext, register
from utils.logger import logger
from utils.timestamps import seconds_to_display


def _install_hint() -> str:
    """الأمر بالمفسّر الذي يشغّل البرنامج فعلًا — لا ‏«pip» أيًّا كان.

    ‏«pip install» يثبّت في أول Python على PATH، وقد لا يكون الذي يشغّل
    الواجهة؛ فيقول المستخدم «مثبَّتة» ويبقى التحذير.
    """
    import sys
    return ("‏python-pptx غير مثبَّتة في مفسّر البرنامج — صيغة الشرائح "
            f'تحتاجها. ثبّتها عبر: "{sys.executable}" -m pip install python-pptx')


INSTALL_HINT = _install_hint()

#: شريحة تتجاوز هذا العدد من النقاط تصير صفحةً تُقرأ لا شريحةً تُعرض.
MAX_BULLETS = 6
MAX_BULLET_CHARS = 120
#: العنوان الطويل يلتفّ على ثلاثة أسطر فيأكل الشريحة
MAX_TITLE_CHARS = 70

_SENTENCE_SPLIT = re.compile(r"(?<=[.!?؟…])\s+")


def _rtl(paragraph, align_right: bool = True) -> None:
    """يضبط ``rtl`` و``algn`` على فقرة داخل الشريحة."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("rtl", "1")
    if align_right:
        p_pr.set("algn", "r")


def _fit(text: str, limit: int) -> str:
    text = " ".join((text or "").split())
    if len(text) <= limit:
        return text
    return text[:limit - 1].rstrip() + "…"


def section_bullets(section: DocumentSection) -> List[str]:
    """نقاط الشريحة: ما كتبته إعادة الصياغة إن وُجد، وإلا يُشتقّ من النصّ.

    ترتيب الأفضلية مقصود. خطةُ الذكاء الاصطناعي تحمل ``bullets`` مكتوبة
    للعرض أصلًا، فهي أولى. وفي المسار الافتراضي (بلا نموذج لغوي) لا
    يوجد إلا كلامٌ منطوق متّصل، فتُؤخذ أوائل الجُمل — وهي في المحاضرة
    عادةً موضع طرح الفكرة قبل شرحها.
    """
    explicit: List[str] = []
    for block in section.blocks:
        if block.kind == "bullets":
            explicit.extend(b for b in block.bullets if (b or "").strip())
    if explicit:
        return [_fit(b, MAX_BULLET_CHARS) for b in explicit[:MAX_BULLETS]]

    sentences: List[str] = []
    for block in section.blocks:
        if block.kind != "paragraph":
            continue
        for sentence in _SENTENCE_SPLIT.split(block.text or ""):
            sentence = sentence.strip()
            # الجملة الأقصر من ثلاثين حرفًا في كلامٍ منطوق شبه دائمًا
            # حشوٌ («طيب»، «تمام يا شباب») لا فكرة.
            if len(sentence) >= 30:
                sentences.append(_fit(sentence, MAX_BULLET_CHARS))
            if len(sentences) >= MAX_BULLETS:
                return sentences
    return sentences


def _speaker_notes(section: DocumentSection) -> str:
    return "\n\n".join(
        (block.text or "").strip()
        for block in section.blocks
        if block.kind == "paragraph" and (block.text or "").strip())


def _add_text_slide(presentation, title: str, lines: List[str],
                    font: str, notes: str = "") -> None:
    from pptx.util import Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = _fit(title, MAX_TITLE_CHARS)
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        _rtl(paragraph)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(30)

    body = slide.placeholders[1].text_frame
    body.clear()
    if not lines:
        lines = ["—"]
    for index, line in enumerate(lines):
        paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
        paragraph.text = line
        paragraph.level = 0
        _rtl(paragraph)
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(18)

    if notes:
        slide.notes_slide.notes_text_frame.text = notes


def _add_picture_slide(presentation, image_path: Path, caption: str,
                       timestamp: Optional[float], font: str,
                       alt: str = "") -> None:
    from pptx.util import Emu, Inches, Pt

    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide_w, slide_h = presentation.slide_width, presentation.slide_height
    margin = Inches(0.45)
    caption_h = Inches(0.75)
    box_w = slide_w - 2 * margin
    box_h = slide_h - 2 * margin - caption_h

    # احتواء بالنسبة الأصلية: التمديد إلى الصندوق يشوّه لقطة الشاشة
    # فيصير النصّ المكتوب عليها غير مقروء — وهو كل قيمتها.
    from PIL import Image
    with Image.open(image_path) as handle:
        width_px, height_px = handle.size
    scale = min(box_w / width_px, box_h / height_px)
    draw_w, draw_h = Emu(int(width_px * scale)), Emu(int(height_px * scale))
    left = Emu(int((slide_w - draw_w) / 2))
    top = Emu(int(margin + (box_h - draw_h) / 2))
    picture = slide.shapes.add_picture(str(image_path), left, top,
                                       width=draw_w, height=draw_h)
    if alt:
        # PowerPoint يعرض «لا يوجد نص بديل» في مدقّق إمكانية الوصول
        # المدمج، تمامًا كـWord. والموضع هنا ``p:cNvPr/@descr``.
        try:
            picture._element._nvXxPr.cNvPr.set("descr", alt)
        except Exception:
            pass

    label = " ".join((caption or "").split())
    if timestamp is not None:
        label = f"{label} — {seconds_to_display(timestamp)}".strip(" —")
    if not label:
        return
    text_box = slide.shapes.add_textbox(
        margin, slide_h - margin - caption_h, box_w, caption_h)
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.text = _fit(label, 140)
    _rtl(paragraph)
    for run in paragraph.runs:
        run.font.name = font
        run.font.size = Pt(13)


def export(ctx: ExportContext) -> Optional[Path]:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ModuleNotFoundError:
        logger.info(INSTALL_HINT)
        return None

    plan = ctx.plan
    if not plan.sections:
        logger.warning("خطة بلا أقسام — تُخطّى الشرائح.")
        return None

    font = getattr(ctx.document_config, "font_family", None) or "Arial"

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)   # ‏16:9
    presentation.slide_height = Inches(7.5)

    facts = [f"المدة: {seconds_to_display(ctx.metadata.duration_seconds)}",
             f"الأقسام: {len(plan.sections)}"]
    if plan.subtitle:
        facts.insert(0, plan.subtitle)
    _add_text_slide(presentation, plan.title or ctx.base_path.stem,
                    facts, font)

    points = [p for p in (plan.key_points or []) if (p or "").strip()]
    if plan.abstract or points:
        lines = points[:MAX_BULLETS] or [_fit(plan.abstract, 220)]
        _add_text_slide(presentation, "الملخّص", lines, font,
                        notes=plan.abstract or "")

    for section in plan.sections:
        _add_text_slide(
            presentation,
            section.title or "قسم",
            ([_fit(section.summary, MAX_BULLET_CHARS)] if section.summary else [])
            + section_bullets(section),
            font,
            notes=_speaker_notes(section))
        for block in section.blocks:
            if block.kind != "figure" or not block.image_filename:
                continue
            image_path = ctx.images_dir / block.image_filename
            if not image_path.is_file():
                continue
            try:
                _add_picture_slide(presentation, image_path, block.caption,
                                   block.timestamp, font,
                                   alt=describe(block, block.image_id))
            except Exception as exc:
                # صورة واحدة تالفة لا تُسقط العرض كلّه
                logger.warning(
                    f"تعذّرت إضافة الشريحة {block.image_filename}: {exc}")

    path = ctx.sibling(".pptx")
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path


register("pptx", export)
