"""قالب الوثيقة الاحترافية: أنماط Word حقيقية، غلاف، فهرس، رأس وتذييل.

لماذا أنماط Word لا تنسيق يدوي لكل فقرة:
    * الفهرس التلقائي يعمل بالأنماط، ولا يعمل بالتنسيق اليدوي.
    * جزء التنقّل في Word يعرض العناوين — فيصبح مستند ساعتين قابلًا
      للتصفّح.
    * تغيير الخط أو اللون لاحقًا يتم من مكان واحد، لا بمرور على آلاف
      الفقرات.
    * حجم الملف يصغر: التنسيق يُعرَّف مرة لا مع كل run.

كل نمط يُضبط على مسار Complex Script أيضًا (``w:cs`` / ``w:szCs``)، وإلا
تجاهله النص العربي — انظر ``document/rtl_utils``.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from document.rtl_utils import (
    add_mixed_text,
    dominant_direction,
    set_paragraph_ltr,
    set_paragraph_rtl,
    set_section_rtl,
    set_table_rtl,
    style_arabic_run,
    style_ltr_run,
)
from utils.timestamps import arabic_datetime

# مسار الشعار الافتراضي داخل المشروع
DEFAULT_LOGO = Path(__file__).resolve().parents[1] / "assets" / "logo.png"


@dataclass(frozen=True)
class Theme:
    """هوية بصرية واحدة للمستند كله.

    الألوان الافتراضية مأخوذة من شعار المشروع (#2D2E82) ليخرج المستند
    متسقًا مع الهوية بدل ألوان عامة.
    """
    font: str = "Arial"
    heading_font: str = "Arial"
    body_pt: float = 12.0
    accent: RGBColor = RGBColor(0x2D, 0x2E, 0x82)      # كحلي الشعار
    accent_soft: RGBColor = RGBColor(0x5C, 0x5D, 0xA8)
    muted: RGBColor = RGBColor(0x6B, 0x72, 0x80)
    rule: str = "D2D4E6"
    accent_hex: str = "2D2E82"
    logo_path: Optional[Path] = None
    logo_width_inches: float = 2.1

    def resolved_logo(self) -> Optional[Path]:
        """يعيد مسار شعار موجود فعلًا، أو ``None``.

        غياب الشعار لا يُفشل التوليد — المستند يخرج بلا شعار فقط.
        """
        candidate = self.logo_path or DEFAULT_LOGO
        return candidate if candidate and Path(candidate).is_file() else None


# أسماء الأنماط المخصّصة
BODY = "نص المستند"
CAPTION = "تسمية شكل"
TIMECODE = "طابع زمني"
LEAD = "مقدمة"


def _set_style_fonts(style, theme: Theme, size_pt: float, *,
                     bold: bool = False, color: Optional[RGBColor] = None,
                     font: Optional[str] = None) -> None:
    """يضبط الخط على مساري Latin و Complex Script معًا."""
    name = font or theme.font
    style.font.name = name
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    for attribute in ("w:ascii", "w:hAnsi", "w:cs"):
        rFonts.set(qn(attribute), name)

    half_points = str(int(round(size_pt * 2)))
    existing = rPr.find(qn("w:szCs"))
    if existing is None:
        szCs = OxmlElement("w:szCs")
        szCs.set(qn("w:val"), half_points)
        rPr.append(szCs)
    else:
        existing.set(qn("w:val"), half_points)
    if bold and rPr.find(qn("w:bCs")) is None:
        rPr.append(OxmlElement("w:bCs"))
    if rPr.find(qn("w:rtl")) is None:
        rPr.append(OxmlElement("w:rtl"))


def _style_paragraph_format(style, *, space_before: float = 0,
                            space_after: float = 6, line_spacing: float = 1.35,
                            keep_with_next: bool = False) -> None:
    paragraph_format = style.paragraph_format
    paragraph_format.space_before = Pt(space_before)
    paragraph_format.space_after = Pt(space_after)
    paragraph_format.line_spacing = line_spacing
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if keep_with_next:
        paragraph_format.keep_with_next = True

    pPr = style.element.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = OxmlElement("w:bidi")
        bidi.set(qn("w:val"), "1")
        pPr.insert(0, bidi)


def build_styles(document: Document, theme: Theme) -> None:
    """يعرّف كل أنماط المستند مرة واحدة."""
    from docx.enum.style import WD_STYLE_TYPE

    normal = document.styles["Normal"]
    _set_style_fonts(normal, theme, theme.body_pt)
    _style_paragraph_format(normal)

    # العناوين المدمجة — يعتمد عليها الفهرس وجزء التنقّل
    for level, size, color in ((1, 17, theme.accent),
                               (2, 14, theme.accent_soft),
                               (3, 12.5, theme.accent_soft)):
        style = document.styles[f"Heading {level}"]
        _set_style_fonts(style, theme, size, bold=True, color=color,
                         font=theme.heading_font)
        _style_paragraph_format(style, space_before=16 if level == 1 else 12,
                                space_after=6, line_spacing=1.2,
                                keep_with_next=True)

    styles = document.styles

    def ensure(name: str):
        try:
            return styles[name]
        except KeyError:
            return styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)

    body = ensure(BODY)
    body.base_style = normal
    _set_style_fonts(body, theme, theme.body_pt)
    _style_paragraph_format(body, space_after=8, line_spacing=1.45)

    lead = ensure(LEAD)
    lead.base_style = normal
    _set_style_fonts(lead, theme, theme.body_pt + 1, color=theme.muted)
    _style_paragraph_format(lead, space_after=10, line_spacing=1.4)

    caption = ensure(CAPTION)
    caption.base_style = normal
    _set_style_fonts(caption, theme, 9.5, color=theme.muted)
    _style_paragraph_format(caption, space_before=3, space_after=12,
                            line_spacing=1.05)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER

    timecode = ensure(TIMECODE)
    timecode.base_style = normal
    _set_style_fonts(timecode, theme, 9.5, bold=True, color=theme.accent_soft)
    _style_paragraph_format(timecode, space_before=10, space_after=2,
                            line_spacing=1.0, keep_with_next=True)


# ---------------------------------------------------------------------
def _field(paragraph, instruction: str, placeholder: str = "",
           theme: Optional[Theme] = None, size_pt: float = 10.0,
           color: Optional[RGBColor] = None):
    """يُدرج حقل Word ديناميكي (PAGE، TOC، …).

    نص العنصر النائب قد يكون عربيًا، فيجب أن يحمل خصائص Complex Script
    مثل أي نص آخر — وإلا خرج بخط الثيم الافتراضي وكسر بوابة الجودة.
    """
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction_element, separate, text, end):
        run._r.append(element)

    if theme is not None:
        if any("؀" <= ch <= "ۿ" for ch in placeholder):
            style_arabic_run(run, theme.font, size_pt, color=color)
        else:
            style_ltr_run(run, theme.font, size_pt, color=color)
    return run


def add_page_number_footer(document: Document, theme: Theme,
                           text: str = "") -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    paragraph.text = ""
    set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.CENTER)
    if text:
        style_arabic_run(paragraph.add_run(f"{text}  —  "), theme.font, 9,
                         color=theme.muted)
    style_arabic_run(paragraph.add_run("صفحة "), theme.font, 9, color=theme.muted)
    _field(paragraph, " PAGE ", "1", theme, 9, theme.muted)
    style_arabic_run(paragraph.add_run(" من "), theme.font, 9, color=theme.muted)
    _field(paragraph, " NUMPAGES ", "1", theme, 9, theme.muted)


def add_header(document: Document, theme: Theme, title: str) -> None:
    """رأس الصفحة: شعار صغير على اليمين وعنوان المستند بجانبه."""
    header = document.sections[0].header
    paragraph = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    paragraph.text = ""
    set_paragraph_rtl(paragraph, WD_ALIGN_PARAGRAPH.RIGHT)

    logo = theme.resolved_logo()
    if logo:
        paragraph.add_run().add_picture(str(logo), width=Inches(0.95))
        style_arabic_run(paragraph.add_run("   "), theme.font, 9)

    if dominant_direction(title) == "rtl":
        add_mixed_text(paragraph, title, theme.font, 9, color=theme.muted)
    else:
        style_ltr_run(paragraph.add_run(title), theme.font, 9,
                      color=theme.muted)
    _horizontal_rule(paragraph, theme.rule, 4)


# ---------------------------------------------------------------------
def _shaded_band(paragraph, hex_color: str) -> None:
    """تظليل خلفية الفقرة — يُستخدم كشريط لوني في الغلاف."""
    pPr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    pPr.append(shading)


def _horizontal_rule(paragraph, color: str, size: int = 8) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    pPr.append(borders)


def _repeat_header_row(row) -> None:
    """يجعل الصف يتكرّر كترويسة عند انقسام الجدول بين الصفحات."""
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def _write_value(paragraph, text: str, theme: Theme, size_pt: float,
                 bold: bool = False, color=None,
                 align=WD_ALIGN_PARAGRAPH.RIGHT) -> None:
    """يكتب قيمة متغيّرة الاتجاه في فقرة، بالطريقة الصحيحة لكل حالة.

    قيمة لاتينية/رقمية خالصة توضع في **فقرة LTR**، فيصح ترتيبها الداخلي
    (``1920×1080`` لا ``1080×1920``) بلا أي محرف تحكّم مضاف.
    """
    if dominant_direction(text) == "rtl":
        set_paragraph_rtl(paragraph, align)
        add_mixed_text(paragraph, text, theme.font, size_pt, bold=bold,
                       color=color)
    else:
        set_paragraph_ltr(paragraph, align)
        style_ltr_run(paragraph.add_run(text), theme.font, size_pt,
                      bold=bold, color=color)


def add_cover_page(document: Document, theme: Theme, *, title: str,
                   subtitle: str, facts: list[tuple[str, str]],
                   kicker: str = "تفريغ ومحتوى مرئي") -> None:
    """غلاف احترافي: شعار، شريط لوني، عنوان، بيانات المصدر."""
    logo = theme.resolved_logo()
    if logo:
        logo_paragraph = document.add_paragraph()
        set_paragraph_rtl(logo_paragraph, WD_ALIGN_PARAGRAPH.CENTER)
        logo_paragraph.paragraph_format.space_before = Pt(30)
        logo_paragraph.paragraph_format.space_after = Pt(22)
        logo_paragraph.add_run().add_picture(
            str(logo), width=Inches(theme.logo_width_inches))
    else:
        for _ in range(2):
            document.add_paragraph()

    band = document.add_paragraph()
    set_paragraph_rtl(band, WD_ALIGN_PARAGRAPH.CENTER)
    band.paragraph_format.space_after = Pt(2)
    _shaded_band(band, theme.accent_hex)
    style_arabic_run(band.add_run(f"  {kicker}  "), theme.font, 10.5,
                     bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))

    document.add_paragraph().paragraph_format.space_after = Pt(6)

    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(4)
    if dominant_direction(title) == "rtl":
        set_paragraph_rtl(heading, WD_ALIGN_PARAGRAPH.CENTER)
        add_mixed_text(heading, title, theme.heading_font, 26, bold=True,
                       color=theme.accent)
    else:
        set_paragraph_ltr(heading, WD_ALIGN_PARAGRAPH.CENTER)
        style_ltr_run(heading.add_run(title), theme.heading_font, 26,
                      bold=True, color=theme.accent)

    if subtitle:
        sub = document.add_paragraph()
        sub.paragraph_format.space_after = Pt(10)
        _write_value(sub, subtitle, theme, 12, color=theme.muted,
                     align=WD_ALIGN_PARAGRAPH.CENTER)

    rule = document.add_paragraph()
    rule.paragraph_format.space_after = Pt(18)
    set_paragraph_rtl(rule, WD_ALIGN_PARAGRAPH.CENTER)
    _horizontal_rule(rule, theme.rule, 12)

    # صف عنوان حقيقي: أنماط Word الملوّنة تعامل الصف الأول كترويسة،
    # فبدونه يظهر أول سطر بيانات ملوّنًا كأنه عنوان.
    table = document.add_table(rows=len(facts) + 1, cols=2)
    table.style = "Light List Accent 1"
    set_table_rtl(table)

    header_cells = (("البيان", theme), ("القيمة", theme))
    for column, (text, _) in enumerate(header_cells):
        paragraph = table.cell(0, column).paragraphs[0]
        set_paragraph_rtl(paragraph)
        style_arabic_run(paragraph.add_run(text), theme.font, 10.5, bold=True)
    _repeat_header_row(table.rows[0])

    for index, (label, value) in enumerate(facts, start=1):
        label_paragraph = table.cell(index, 0).paragraphs[0]
        set_paragraph_rtl(label_paragraph)
        style_arabic_run(label_paragraph.add_run(label), theme.font, 10.5,
                         bold=True, color=theme.accent)
        _write_value(table.cell(index, 1).paragraphs[0], value, theme, 10.5)

    document.add_paragraph()
    generated = document.add_paragraph()
    set_paragraph_rtl(generated, WD_ALIGN_PARAGRAPH.CENTER)
    # صيغة عربية طبيعية: لا مجموعات أرقام مفصولة بشرطات تنعكس بصريًا
    add_mixed_text(generated, f"أُنشئ آليًا في {arabic_datetime()}",
                   theme.font, 9, color=theme.muted)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table_of_contents(document: Document, theme: Theme,
                          heading: str = "المحتويات") -> None:
    """فهرس تلقائي يتحدّث داخل Word بـ F9 أو عند الطباعة."""
    title = document.add_paragraph()
    set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.RIGHT)
    style_arabic_run(title.add_run(heading), theme.heading_font, 17,
                     bold=True, color=theme.accent)

    holder = document.add_paragraph()
    set_paragraph_rtl(holder)
    _field(holder, ' TOC \\o "1-3" \\h \\z \\u ',
           "اضغط داخل هذا الإطار ثم F9 لتحديث الفهرس",
           theme, theme.body_pt, theme.muted)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_figure_index(document: Document, theme: Theme,
                     heading: str = "فهرس الأشكال") -> None:
    """قائمة بكل الأشكال وتوقيتاتها، تُبنى من حقول SEQ في التسميات.

    لمراجعة محاضرة، هذه القائمة هي أسرع طريق إلى اللقطة المطلوبة: تُقرأ
    كفهرس بصري للفيديو كله. تتحدّث داخل Word بـ F9 مثل جدول المحتويات.

    المعرّف ``Figure`` لاتيني عمدًا رغم أن التسمية المعروضة عربية:
    معرّفات SEQ العربية تعمل في Word لكنها تنكسر عند فتح الملف في
    محرّرات أخرى، والمعرّف لا يظهر للقارئ أصلًا.
    """
    title = document.add_paragraph()
    set_paragraph_rtl(title, WD_ALIGN_PARAGRAPH.RIGHT)
    style_arabic_run(title.add_run(heading), theme.heading_font, 14,
                     bold=True, color=theme.accent)

    holder = document.add_paragraph()
    set_paragraph_rtl(holder)
    _field(holder, ' TOC \\h \\z \\c "Figure" ',
           "اضغط داخل هذا الإطار ثم F9 لتحديث فهرس الأشكال",
           theme, theme.body_pt, theme.muted)

    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_sequence_number(paragraph, identifier: str, theme: Theme,
                        size_pt: float = 9.5, fallback: int = 1,
                        color: Optional[RGBColor] = None):
    """رقم شكل محسوب بحقل ``SEQ`` بدل رقم نصّي ثابت.

    الرقم النصّي يبدو صحيحًا لكنه لا يُبنى منه فهرس أشكال، ويصير خاطئًا
    لحظة أن يحذف المستخدم صورة من المستند. حقل SEQ يعيد الترقيم تلقائيًا
    ويغذّي ``add_figure_index``.
    """
    return _field(paragraph, f" SEQ {identifier} \\* ARABIC ",
                  str(fallback), theme, size_pt, color)


def _modernize_compatibility(document: Document) -> None:
    """يرفع إعدادات توافق الإصدار عن قالب python-docx الافتراضي.

    قالب python-docx المدمج (المبني على default.docx القديم) يحمل في
    ``word/settings.xml`` عنصر ``w:compat`` بقيمتين تُخرِّبان مستندات
    RTL تحديدًا:

    * ``compatibilityMode = 14`` (Word 2010) — يفتح Word الحديث الملف
      في «وضع التوافق» (Compatibility Mode) ويستخدم محرك التنسيق
      والعدالة (justification) القديم لنصوص Complex Script، فتظهر
      محاذاة الفقرات مكسورة بصريًا رغم أن ``w:jc``/``w:bidi`` صحيحان
      تمامًا في XML. رفعها إلى 15 (Word 2013+) يزيل الشارة ويشغّل
      محرك التخطيط الحديث.
    * ``doNotFlipMirrorIndents = 1`` — يمنع Word من عكس المسافات
      البادئة (``w:ind``) للفقرات ثنائية الاتجاه، فأي مسافة بادئة
      محسوبة بمنطق LTR (تعداد نقطي، اقتباس بمسافة بادئة يسراً...)
      تبقى ملتصقة بالهامش الأيسر بدل أن تنتقل تلقائيًا لليمين. هذا هو
      سبب ظهور القوائم والاقتباسات منزاحة عن الهامش الأيمن رغم صحة
      محاذاة النص نفسه. تعطيلها (0) يعيد لـ Word سلوكه الطبيعي في
      عكس المسافات البادئة لفقرات RTL.

    يُستدعى مرة واحدة عقب ``Document()`` (من ``configure_page``)، قبل
    أي فقرة أو نمط.
    """
    compat = document.settings.element.find(qn("w:compat"))
    if compat is None:
        return
    for setting in compat.findall(qn("w:compatSetting")):
        name = setting.get(qn("w:name"))
        if name == "compatibilityMode":
            setting.set(qn("w:val"), "15")
        elif name == "doNotFlipMirrorIndents":
            setting.set(qn("w:val"), "0")


def configure_page(document: Document, theme: Theme) -> None:
    _modernize_compatibility(document)
    section = document.sections[0]
    set_section_rtl(section)
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
