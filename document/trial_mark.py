"""علامة «نسخة تجريبية» المائية على مخرجات باقة «تجربة».

الباقة تُعرَّف من العقد الموقَّع (حقل ``tier`` أو خطة تبدأ بـ«تجربة»)،
فلا تُرفع العلامة بتغيير إعداد محلي. تُطبَّق في نقطتين مركزيتين لا في
كل مُصيِّر: ``DocumentGenerator.generate`` (مستند Word، ومنه يتولّد
PDF) و``write_exports`` (HTML وPPTX وحزم SCORM).

العلامة المائية الأساسية لاتينية («MAEEN · TRIAL») لأن WordArt لا يشكّل
العربية بثبات، ويُضاف إليها سطر عربي في ترويسة كل صفحة.
"""
from __future__ import annotations

import os
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Optional

from utils.logger import logger

MARK_LATIN = "MAEEN · TRIAL"
MARK_ARABIC = "نسخة تجريبية من «معين» — فعّل باقة مدفوعة لمستند بلا علامة مائية"
_HTML_FLAG = "data-maeen-trial"


def is_trial() -> bool:
    """هل الترخيص الحالي من باقة «تجربة»؟ لا يرمي ولا يلمس الشبكة."""
    try:
        from licensing import app_gate

        return app_gate.is_trial()
    except Exception:                                   # noqa: BLE001
        return False


# ── Word ──────────────────────────────────────────────────────────────
_VML_NS = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
           'xmlns:v="urn:schemas-microsoft-com:vml" '
           'xmlns:o="urn:schemas-microsoft-com:office:office" '
           'xmlns:w10="urn:schemas-microsoft-com:office:word"')

_WATERMARK_XML = f"""<w:p {_VML_NS}>
<w:pPr><w:pStyle w:val="Header"/><w:jc w:val="center"/></w:pPr>
<w:r><w:rPr><w:noProof/></w:rPr><w:pict>
<v:shapetype id="_x0000_t136" coordsize="21600,21600" o:spt="136" adj="10800"
 path="m@7,l@8,m@5,21600l@6,21600e">
<v:formulas><v:f eqn="sum #0 0 10800"/><v:f eqn="prod #0 2 1"/><v:f eqn="sum 21600 0 @1"/>
<v:f eqn="sum 0 0 @2"/><v:f eqn="sum 21600 0 @3"/><v:f eqn="if @0 @3 0"/>
<v:f eqn="if @0 21600 @1"/><v:f eqn="if @0 0 @2"/><v:f eqn="if @0 @4 21600"/>
<v:f eqn="mid @5 @6"/><v:f eqn="mid @8 @5"/><v:f eqn="mid @7 @8"/>
<v:f eqn="mid @6 @7"/><v:f eqn="sum @6 0 @5"/></v:formulas>
<v:path textpathok="t" o:connecttype="custom" o:connectlocs="@9,0;@10,10800;@11,21600;@12,10800"
 o:connectangles="270,180,90,0"/>
<v:textpath on="t" fitshape="t"/>
<v:handles><v:h position="#0,bottomRight" xrange="6629,14971"/></v:handles>
<o:lock v:ext="edit" text="t" shapetype="t"/></v:shapetype>
<v:shape id="PowerPlusWaterMarkObject{{n}}" o:spid="_x0000_s{{spid}}" type="#_x0000_t136"
 style="position:absolute;margin-left:0;margin-top:0;width:470pt;height:94pt;rotation:315;z-index:-251658752;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin"
 o:allowincell="f" fillcolor="#c8c8c8" stroked="f">
<v:fill opacity=".45"/>
<v:textpath style="font-family:&quot;Arial&quot;;font-size:1pt" string="{MARK_LATIN}"/>
<w10:wrap anchorx="margin" anchory="margin"/>
</v:shape></w:pict></w:r></w:p>"""


def _add_to_header(header, counter: int, force: bool) -> None:
    from lxml import etree

    if header.is_linked_to_previous and not force:
        return                                  # يرثها من القسم الأول
    element = etree.fromstring(
        _WATERMARK_XML.replace("{n}", str(counter + 1))
        .replace("{spid}", str(2049 + counter)))
    header._element.insert(0, element)
    # سطر عربي مرئي في الترويسة
    paragraph = header.add_paragraph()
    paragraph.alignment = 1                     # وسط
    run = paragraph.add_run(MARK_ARABIC)
    run.font.size = _pt(8)
    run.font.color.rgb = _rgb(0xB0, 0x30, 0x30)


def _pt(value):
    from docx.shared import Pt

    return Pt(value)


def _rgb(r, g, b):
    from docx.shared import RGBColor

    return RGBColor(r, g, b)


def mark_docx_object(document) -> bool:
    """يضيف العلامة إلى كل أقسام مستند python-docx. يعيد ``True`` إن طُبّقت."""
    if not is_trial():
        return False
    try:
        for index, section in enumerate(document.sections):
            _add_to_header(section.header, index, force=index == 0)
            if section.different_first_page_header_footer:
                _add_to_header(section.first_page_header, index + 100,
                               force=index == 0)
        return True
    except Exception as exc:                            # noqa: BLE001
        logger.warning(f"تعذّر وضع العلامة المائية على Word: {exc}")
        return False


# ── HTML ──────────────────────────────────────────────────────────────
_HTML_SNIPPET = (
    f'<div {_HTML_FLAG}="1" style="position:fixed;inset:0;pointer-events:none;z-index:2147483647;'
    'display:flex;align-items:center;justify-content:center;overflow:hidden">'
    f'<span style="transform:rotate(-30deg);font:700 min(14vw,120px)/1 Arial,sans-serif;'
    f'color:rgba(120,120,120,.18);white-space:nowrap">{MARK_LATIN}</span></div>'
    f'<div {_HTML_FLAG}="1" style="position:fixed;top:0;left:0;right:0;background:#b03030;color:#fff;'
    'font:12px/1.6 system-ui,sans-serif;text-align:center;padding:3px 8px;z-index:2147483647" dir="rtl">'
    f'{MARK_ARABIC}</div>')


def mark_html_text(text: str) -> str:
    if _HTML_FLAG in text:
        return text
    lowered = text.lower()
    index = lowered.rfind("</body>")
    if index < 0:
        return text + _HTML_SNIPPET
    return text[:index] + _HTML_SNIPPET + text[index:]


def _mark_html_file(path: Path) -> None:
    text = path.read_text(encoding="utf-8")
    path.write_text(mark_html_text(text), encoding="utf-8")


# ── PowerPoint ────────────────────────────────────────────────────────
def _mark_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    presentation = Presentation(str(path))
    width, height = presentation.slide_width, presentation.slide_height
    for slide in presentation.slides:
        box = slide.shapes.add_textbox(int(width * 0.1), int(height * 0.38),
                                       int(width * 0.8), int(height * 0.24))
        box.rotation = -25
        frame = box.text_frame
        frame.text = MARK_LATIN
        paragraph = frame.paragraphs[0]
        paragraph.alignment = 2                         # وسط
        run = paragraph.runs[0]
        run.font.size = Pt(54)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xC8, 0xC8, 0xC8)
        # خلف المحتوى لا فوقه
        tree = slide.shapes._spTree
        tree.remove(box._element)
        tree.insert(2, box._element)
    presentation.save(str(path))


# ── SCORM / أي zip فيه صفحات HTML ─────────────────────────────────────
def _mark_zip(path: Path) -> None:
    fd, tmp_name = tempfile.mkstemp(dir=str(path.parent), suffix=".zip.tmp")
    os.close(fd)
    tmp = Path(tmp_name)
    try:
        with zipfile.ZipFile(path) as source, \
                zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename.lower().endswith((".html", ".htm")):
                    data = mark_html_text(data.decode("utf-8")).encode("utf-8")
                target.writestr(item, data)
        shutil.move(str(tmp), str(path))
    finally:
        tmp.unlink(missing_ok=True)


def mark_file(path: Optional[Path]) -> bool:
    """يضع العلامة على ملف مُصيَّر حسب امتداده. لا يرمي أبدًا."""
    if path is None or not is_trial():
        return False
    path = Path(path)
    suffix = path.suffix.lower()
    try:
        if suffix in (".html", ".htm"):
            _mark_html_file(path)
        elif suffix == ".pptx":
            _mark_pptx(path)
        elif suffix == ".zip":
            _mark_zip(path)
        else:
            return False                                # PDF يأتي من Word المعلَّم
        return True
    except Exception as exc:                            # noqa: BLE001
        logger.warning(f"تعذّر وضع العلامة المائية على {path.name}: {exc}")
        return False
