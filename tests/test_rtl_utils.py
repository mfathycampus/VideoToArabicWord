"""تحقق من صحة الـ XML المولّد مقابل ترتيب مخطط OOXML."""
import sys, re
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from docx import Document
from docx.shared import RGBColor
from docx.oxml.ns import qn

from document.rtl_utils import (
    set_paragraph_rtl, style_arabic_run, set_table_rtl,
    set_section_rtl, configure_document_defaults,
)

# تسلسل المخطط الرسمي (ECMA-376) للعناصر المستخدمة
PPR_SEQ = ["w:pStyle","w:keepNext","w:keepLines","w:pageBreakBefore","w:framePr",
    "w:widowControl","w:numPr","w:suppressLineNumbers","w:pBdr","w:shd","w:tabs",
    "w:suppressAutoHyphens","w:kinsoku","w:wordWrap","w:overflowPunct","w:topLinePunct",
    "w:autoSpaceDE","w:autoSpaceDN","w:bidi","w:adjustRightInd","w:snapToGrid",
    "w:spacing","w:ind","w:contextualSpacing","w:mirrorIndents","w:suppressOverlap",
    "w:jc","w:textDirection","w:textAlignment","w:textboxTightWrap","w:outlineLvl",
    "w:divId","w:cnfStyle","w:rPr","w:sectPr","w:pPrChange"]

RPR_SEQ = ["w:rStyle","w:rFonts","w:b","w:bCs","w:i","w:iCs","w:caps","w:smallCaps",
    "w:strike","w:dstrike","w:outline","w:shadow","w:emboss","w:imprint","w:noProof",
    "w:snapToGrid","w:vanish","w:webHidden","w:color","w:spacing","w:w","w:kern",
    "w:position","w:sz","w:szCs","w:highlight","w:u","w:effect","w:bdr","w:shd",
    "w:fitText","w:vertAlign","w:rtl","w:cs","w:em","w:lang","w:eastAsianLayout",
    "w:specVanish","w:oMath"]

TBLPR_SEQ = ["w:tblStyle","w:tblpPr","w:tblOverlap","w:bidiVisual",
    "w:tblStyleRowBandSize","w:tblStyleColBandSize","w:tblW","w:jc",
    "w:tblCellSpacing","w:tblInd","w:tblBorders","w:shd","w:tblLayout",
    "w:tblCellMar","w:tblLook","w:tblCaption","w:tblDescription","w:tblPrChange"]


def _local_tags(element):
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    return ["w:" + child.tag.replace(ns, "") for child in element]


def _assert_ordered(element, sequence, label):
    tags = _local_tags(element)
    indices = [sequence.index(t) for t in tags if t in sequence]
    assert indices == sorted(indices), (
        f"{label}: ترتيب مخالف للمخطط -> {tags}"
    )


@pytest.fixture
def doc():
    d = Document()
    configure_document_defaults(d, "Arial", 13)
    return d


def test_paragraph_bidi_precedes_jc(doc):
    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    pPr = p._p.get_or_add_pPr()
    tags = _local_tags(pPr)
    assert "w:bidi" in tags, "عنصر w:bidi مفقود"
    assert tags.index("w:bidi") < tags.index("w:jc"), "w:bidi يجب أن يسبق w:jc"
    _assert_ordered(pPr, PPR_SEQ, "pPr")


def test_run_carries_complex_script_properties(doc):
    p = doc.add_paragraph()
    set_paragraph_rtl(p)
    r = p.add_run("الذكاء الاصطناعي")
    style_arabic_run(r, "Amiri", 13, bold=True)
    rPr = r._r.get_or_add_rPr()

    # الخط العربي يُقرأ من w:cs حصريًا
    assert rPr.rFonts.get(qn("w:cs")) == "Amiri", "خط Complex Script غير مضبوط"
    # الحجم العربي يُقرأ من w:szCs حصريًا (بنصف النقطة)
    assert rPr.find(qn("w:szCs")).get(qn("w:val")) == "26", "حجم Complex Script غير مضبوط"
    # الغامق العربي يُقرأ من w:bCs
    assert rPr.find(qn("w:bCs")) is not None, "w:bCs مفقود"
    # اتجاه الـ run
    assert rPr.find(qn("w:rtl")) is not None, "w:rtl مفقود"
    _assert_ordered(rPr, RPR_SEQ, "rPr")


def test_table_uses_bidivisual(doc):
    t = doc.add_table(rows=2, cols=2)
    set_table_rtl(t)
    tblPr = t._tbl.tblPr
    tags = _local_tags(tblPr)
    assert "w:bidiVisual" in tags, "الجداول تحتاج w:bidiVisual وليس w:bidi"
    assert "w:bidi" not in tags, "w:bidi غير صالح داخل tblPr"
    _assert_ordered(tblPr, TBLPR_SEQ, "tblPr")


def test_section_is_rtl(doc):
    sectPr = doc.sections[0]._sectPr
    assert sectPr.find(qn("w:bidi")) is not None, "اتجاه القسم غير مضبوط"


def test_idempotent(doc):
    """استدعاء الدوال مرتين لا يكرر العناصر."""
    p = doc.add_paragraph()
    set_paragraph_rtl(p); set_paragraph_rtl(p)
    pPr = p._p.get_or_add_pPr()
    assert _local_tags(pPr).count("w:bidi") == 1

    r = p.add_run("نص")
    style_arabic_run(r, "Arial", 13); style_arabic_run(r, "Arial", 14)
    rPr = r._r.get_or_add_rPr()
    assert _local_tags(rPr).count("w:szCs") == 1
    assert rPr.find(qn("w:szCs")).get(qn("w:val")) == "28"


def test_document_roundtrips(doc, tmp_path):
    p = doc.add_paragraph(); set_paragraph_rtl(p)
    style_arabic_run(p.add_run("مرحبا بالعالم"), "Arial", 13, color=RGBColor(0,0,0))
    out = tmp_path / "rtl.docx"
    doc.save(str(out))
    reopened = Document(str(out))
    assert "مرحبا بالعالم" in reopened.paragraphs[-1].text


def test_right_aligned_arabic_is_written_as_word_reads_it():
    """‏Word يقرأ jc منطقيًّا في فقرة bidi: left = يمين الصفحة.

    كُتب ``right`` سنواتٍ فظهر كل نص عربي على اليسار — ولم يفشل اختبار،
    لأن الاختبارات فحصت ما كُتب لا ما يُعرض.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    from document.rtl_utils import (
        finalize_bidi_alignment,
        set_paragraph_ltr,
        set_paragraph_rtl,
    )

    document = Document()
    arabic = document.add_paragraph("نص عربي")
    set_paragraph_rtl(arabic, WD_ALIGN_PARAGRAPH.RIGHT)
    centered = document.add_paragraph("وسط")
    set_paragraph_rtl(centered, WD_ALIGN_PARAGRAPH.CENTER)
    latin = document.add_paragraph("1920x1080")
    set_paragraph_ltr(latin, WD_ALIGN_PARAGRAPH.RIGHT)

    finalize_bidi_alignment(document)

    jc = lambda p: p._p.pPr.find(qn("w:jc")).get(qn("w:val"))  # noqa: E731
    assert jc(arabic) == "left"
    assert jc(centered) == "center"
    assert jc(latin) == "right"          # فقرة LTR: المحاذاة مرئية كما هي
