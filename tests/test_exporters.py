"""بوابات المُصيِّرات الإضافية.

ما تحرسه هذه الاختبارات ليس شكل المخرج بل **عقد الفشل**: مخرجات
إضافية لا شروط نجاح. مُصيِّر يسقط — أو اعتمادٌ خارجي يغيب — يجب أن
يُتخطّى وحده وتُكتب البقية. الانحدار هنا صامت بطبعه: لا يظهر على جهاز
المطوّر الذي يملك كل الاعتمادات، ويظهر عند المعلّم وحده.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
from PIL import Image

from config.schemas import (
    DocumentBlock,
    DocumentPlan,
    DocumentSection,
    VideoMetadata,
)
from config.settings import DocumentConfig
from document.exporters import ExportContext, available_formats, write_exports


# ---------------------------------------------------------------------
# تجهيزات
# ---------------------------------------------------------------------
def _plan() -> DocumentPlan:
    return DocumentPlan(
        title="مقدّمة في الشبكات",
        subtitle="محاضرة تجريبية",
        abstract="نظرة عامة على طبقات الشبكة.",
        key_points=["الطبقة المادية", "طبقة النقل"],
        sections=[
            DocumentSection(
                title="القسم الأول — التعريف", start_timestamp=0.0,
                summary="ما الشبكة ولماذا نحتاجها.",
                blocks=[
                    DocumentBlock(kind="paragraph", segment_ids=[1],
                                  text="الشبكة مجموعة أجهزة متصلة تتبادل "
                                       "البيانات وفق بروتوكول متفق عليه."),
                    DocumentBlock(kind="figure", image_id=1,
                                  image_filename="f1.jpg", timestamp=12.0,
                                  caption="شكل 1 — طبقات OSI",
                                  ocr_text="OSI Model"),
                ]),
            DocumentSection(
                title="القسم الثاني — البروتوكولات", start_timestamp=90.0,
                blocks=[
                    DocumentBlock(kind="paragraph", segment_ids=[2],
                                  text="بروتوكول TCP يضمن الوصول والترتيب، "
                                       "بينما UDP يتخلّى عنهما مقابل السرعة."),
                    DocumentBlock(kind="bullets",
                                  bullets=["TCP موثوق", "UDP سريع"]),
                ]),
            DocumentSection(
                title="القسم الثالث — الخلاصة", start_timestamp=200.0,
                blocks=[DocumentBlock(kind="paragraph", segment_ids=[3],
                                      text="ملخّص ما سبق في جملة واحدة "
                                           "تكفي للمراجعة السريعة.")]),
        ],
    )


def _metadata(tmp_path: Path) -> VideoMetadata:
    return VideoMetadata(
        filename="lecture.mp4", path=tmp_path / "lecture.mp4",
        duration_seconds=300.0, width=1920, height=1080, fps=30.0,
        codec="h264", has_audio=True, has_video=True)


@pytest.fixture
def ctx(tmp_path: Path) -> ExportContext:
    images = tmp_path / "keyframes"
    images.mkdir()
    Image.new("RGB", (640, 360), (40, 60, 120)).save(images / "f1.jpg")
    return ExportContext(
        plan=_plan(), metadata=_metadata(tmp_path), images_dir=images,
        base_path=tmp_path / "lecture.docx",
        document_config=DocumentConfig())


# ---------------------------------------------------------------------
# السجلّ وعقد الفشل
# ---------------------------------------------------------------------
def test_registry_exposes_the_four_builtin_formats():
    assert set(available_formats()) >= {"html", "pdf", "pptx", "chapters"}


def test_unknown_format_is_skipped_not_raised(ctx):
    assert write_exports(ctx, ["صيغة-لا-وجود-لها"]) == []


def test_failing_exporter_does_not_stop_the_others(ctx, monkeypatch):
    """الحارس الأهم في الملف كلّه.

    بلا هذا العزل يصير أي عطل في مُصيِّر تجريبي عطلًا في المهمة كاملة —
    أي أن إضافة صيغة جديدة تُخاطر بمستند Word نفسه، وهو عكس الغرض من
    فصل المُصيِّرات ابتداءً.
    """
    import document.exporters as exporters

    def boom(_ctx):
        raise RuntimeError("عطل مفتعل")

    monkeypatch.setitem(exporters._REGISTRY, "قنبلة", boom)
    written = write_exports(ctx, ["قنبلة", "chapters"])
    assert [p.suffix for p in written] == [".txt"]


def test_duplicate_formats_run_once(ctx):
    assert len(write_exports(ctx, ["chapters", "chapters", "CHAPTERS"])) == 1


# ---------------------------------------------------------------------
# HTML
# ---------------------------------------------------------------------
def test_html_is_one_self_contained_file(ctx):
    from document.exporters.html_export import export

    path = export(ctx)
    text = path.read_text(encoding="utf-8")

    assert path.name == "lecture.html"
    # بلا أي مرجع خارجي: الملف يُفتح على جهاز بلا إنترنت
    assert "http://" not in text and "https://" not in text
    assert "<script src=" not in text and "<link " not in text
    assert 'data:image/jpeg;base64,' in text      # الصورة مضمّنة
    assert 'dir="rtl"' in text


def test_html_escapes_markup_in_content(tmp_path):
    """نصّ التفريغ مصدرٌ غير موثوق شكليًّا.

    لقطة شاشة لمحرّر كود تُنتج OCR فيه ``<script>``، وWhisper يُفرّغ
    «أصغر من» نطقًا. بلا هروبٍ صارم تنكسر الصفحة — أو تُنفّذ.
    """
    from document.exporters.html_export import render_html

    plan = _plan()
    plan.sections[0].blocks[0].text = '<script>alert("x")</script> و & و <b>'
    images = tmp_path / "keyframes"
    images.mkdir()
    ctx = ExportContext(plan=plan, metadata=_metadata(tmp_path),
                        images_dir=images, base_path=tmp_path / "l.docx",
                        document_config=DocumentConfig())
    text = render_html(ctx)
    assert "<script>alert" not in text
    assert "&lt;script&gt;" in text


def test_html_skips_missing_image_without_breaking(tmp_path):
    from document.exporters.html_export import render_html

    images = tmp_path / "keyframes"
    images.mkdir()                      # المجلد موجود والصورة مفقودة
    ctx = ExportContext(plan=_plan(), metadata=_metadata(tmp_path),
                        images_dir=images, base_path=tmp_path / "l.docx",
                        document_config=DocumentConfig())
    text = render_html(ctx)
    assert "<img" not in text
    assert "القسم الأول" in text        # النصّ كامل رغم سقوط الشكل


def test_html_links_images_when_they_exceed_the_embed_cap(ctx, monkeypatch):
    import document.exporters.html_export as html_export

    monkeypatch.setattr(html_export, "MAX_EMBEDDED_IMAGE_BYTES", 10)
    text = html_export.render_html(ctx)
    assert 'src="keyframes/f1.jpg"' in text
    assert "base64" not in text


def test_html_section_count_matches_plan(ctx):
    from document.exporters.html_export import render_html

    text = render_html(ctx)
    assert text.count('class="chapter"') == len(ctx.plan.sections)


# ---------------------------------------------------------------------
# الفصول
# ---------------------------------------------------------------------
def test_chapters_follow_youtube_rules(ctx):
    from document.exporters.chapters_export import export

    path = export(ctx)
    lines = [ln for ln in path.read_text(encoding="utf-8").splitlines()
             if ln and not ln.startswith("#")]

    assert len(lines) >= 3                 # الحد الأدنى الذي يقبله يوتيوب
    assert lines[0].startswith("0:00")     # أول فصل عند الصفر بالضبط


def test_chapters_merge_sections_shorter_than_the_minimum():
    from document.exporters.chapters_export import build_chapters

    plan = _plan()
    plan.sections[1].start_timestamp = 3.0      # ثلاث ثوانٍ بعد الأول
    chapters = build_chapters(plan, duration_seconds=300.0)
    assert [start for start, _ in chapters] == [0.0, 200.0] or chapters == []


def test_chapters_refuses_to_write_a_list_youtube_would_reject(tmp_path):
    """قائمة أقصر من ثلاثة فصول تُلغى بدل كتابتها.

    يوتيوب يرفض القائمة **كلّها** عندها، فملفٌ مكتوبٌ هنا يعني مستخدمًا
    يلصق ولا يحدث شيء ولا رسالة تقول لماذا.
    """
    from document.exporters.chapters_export import export

    plan = _plan()
    plan.sections = plan.sections[:2]
    ctx = ExportContext(plan=plan, metadata=_metadata(tmp_path),
                        images_dir=tmp_path, base_path=tmp_path / "l.docx",
                        document_config=DocumentConfig())
    assert export(ctx) is None
    assert not (tmp_path / "l.chapters.txt").exists()


# ---------------------------------------------------------------------
# PDF — الاعتماد الخارجي
# ---------------------------------------------------------------------
def test_pdf_is_skipped_when_libreoffice_is_absent(ctx, monkeypatch, tmp_path):
    """جهاز المعلّم بلا LibreOffice — وهو الحال الافتراضي."""
    import document.exporters.pdf_export as pdf_export

    monkeypatch.setattr(pdf_export, "find_soffice", lambda: None)
    ctx.docx_path = tmp_path / "lecture.docx"
    ctx.docx_path.write_bytes(b"PK\x03\x04")
    assert pdf_export.export(ctx) is None


def test_pdf_is_skipped_when_there_is_no_document(ctx):
    """مسار ``--transcript-only``: لا مستند Word أصلًا ليُحوَّل."""
    from document.exporters.pdf_export import export

    assert ctx.docx_path is None
    assert export(ctx) is None


def test_pdf_reports_failure_when_soffice_exits_zero_without_a_file(
        ctx, monkeypatch, tmp_path):
    """‏soffice يعيد صفرًا حتى حين يفشل — وجود الملف هو المحك الوحيد."""
    import subprocess

    import document.exporters.pdf_export as pdf_export

    monkeypatch.setattr(pdf_export, "find_soffice", lambda: Path("/bin/true"))
    monkeypatch.setattr(
        subprocess, "run",
        lambda *a, **k: subprocess.CompletedProcess(a, 0, "", ""))
    ctx.docx_path = tmp_path / "lecture.docx"
    ctx.docx_path.write_bytes(b"PK\x03\x04")
    assert pdf_export.export(ctx) is None


@pytest.mark.skipif(
    __import__("document.exporters.pdf_export", fromlist=["x"]).find_soffice()
    is None, reason="LibreOffice غير مثبَّت على هذا الجهاز")
def test_pdf_leaves_no_libreoffice_debris_in_the_output_folder(ctx, tmp_path):
    """انحدارٌ قِيس فعلًا على تحويل حقيقي.

    التحويل مباشرةً إلى مجلد المهمة خلّف فيه ``.~lock.<الاسم>#`` وملفًّا
    ``.tmp`` بحجم الـ PDF كاملًا. الأول مخفيّ على ويندوز فلا يراه
    المستخدم ولا يحذفه أحد، والثاني يضاعف حجم مخرجات كل محاضرة.
    """
    from docx import Document

    from document.exporters.pdf_export import convert

    source = tmp_path / "مستند.docx"
    document = Document()
    document.add_paragraph("محتوى تجريبي للتحويل.")
    document.save(str(source))

    out = tmp_path / "out"
    produced = convert(source, out)
    assert produced is not None and produced.is_file()
    assert [f.name for f in out.iterdir()] == ["مستند.pdf"]


# ---------------------------------------------------------------------
# الشرائح
# ---------------------------------------------------------------------
def test_pptx_bullets_prefer_explicit_ones_over_derived_sentences():
    from document.exporters.pptx_export import section_bullets

    plan = _plan()
    assert section_bullets(plan.sections[1]) == ["TCP موثوق", "UDP سريع"]


def test_pptx_derives_bullets_from_prose_when_none_are_written():
    from document.exporters.pptx_export import section_bullets

    bullets = section_bullets(_plan().sections[0])
    assert bullets and all(len(b) <= 120 for b in bullets)


def test_pptx_builds_a_deck_with_a_slide_per_section_and_figure(ctx):
    pptx = pytest.importorskip("pptx", reason="python-pptx اختيارية")
    from document.exporters.pptx_export import export

    path = export(ctx)
    assert path is not None and path.suffix == ".pptx"
    deck = pptx.Presentation(str(path))
    # غلاف + ملخّص + ثلاثة أقسام + شكل واحد
    assert len(deck.slides.__iter__.__self__._sldIdLst) == 6


def test_pptx_is_skipped_when_the_library_is_absent(ctx, monkeypatch):
    import builtins

    real_import = builtins.__import__

    def blocked(name, *args, **kwargs):
        if name == "pptx" or name.startswith("pptx."):
            raise ModuleNotFoundError("No module named 'pptx'")
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", blocked)
    from document.exporters.pptx_export import export

    assert export(ctx) is None


def test_pdf_falls_back_to_word_when_libreoffice_is_missing(monkeypatch, tmp_path):
    """‏python-pptx وLibreOffice غائبان عند المستخدم؛ Word موجود."""
    from document.exporters import pdf_export

    source = tmp_path / "درس.docx"
    source.write_bytes(b"docx")
    calls = {}

    def fake_word(docx, target):
        calls["args"] = (docx, target)
        Path(target).write_bytes(b"%PDF")
        return Path(target)

    monkeypatch.setattr(pdf_export, "find_soffice", lambda: None)
    monkeypatch.setattr(pdf_export, "convert_with_word", fake_word)
    produced = pdf_export.convert(source, tmp_path / "out")
    assert produced == tmp_path / "out" / "درس.pdf"
    assert calls["args"][0] == source


def test_word_conversion_passes_paths_through_the_environment(monkeypatch, tmp_path):
    """أسماء عربية وعلامات اقتباس لا تُحقن في نصّ السكربت."""
    import subprocess
    import sys

    from document.exporters import pdf_export

    monkeypatch.setattr(sys, "platform", "win32")
    monkeypatch.setattr(pdf_export.shutil, "which", lambda name: "powershell")
    monkeypatch.setattr(subprocess, "CREATE_NO_WINDOW", 0, raising=False)
    target = tmp_path / "درس 'أ'.pdf"
    seen = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"], seen["env"] = cmd, kwargs["env"]
        target.write_bytes(b"%PDF")
        return subprocess.CompletedProcess(cmd, 0, "", "")

    monkeypatch.setattr(pdf_export.subprocess, "run", fake_run)
    assert pdf_export.convert_with_word(tmp_path / "درس 'أ'.docx", target) == target
    assert "درس" not in " ".join(seen["cmd"])
    assert seen["env"]["VTAD_PDF"].endswith("درس 'أ'.pdf")
