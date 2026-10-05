"""الحزمة الثالثة: OCR باللغات المثبَّتة، وفصول لا تُحذف بصمت، وفحص بصري للشرائح.

الفحص البصري هنا هو ما كنت أفعله يدويًّا (تحويل إلى صور وعين ترى فيضان
النصّ): يحوّل الشرائح عبر LibreOffice إلى PDF ثم يقرأ مواضع الكلمات.
"""
import logging
import shutil
import subprocess
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest

from config.schemas import DocumentBlock, DocumentPlan, DocumentSection


# ---------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------
@pytest.fixture
def ocr(monkeypatch):
    from video import ocr as module

    monkeypatch.setattr(module, "_cached_langs", None)
    return module


@pytest.mark.parametrize("installed, expected", [
    (["ara", "eng", "osd"], "ara+eng"),
    (["eng", "osd"], "eng"),          # بلا عربية: تُقرأ الإنجليزية بدل الفشل
    (["ara"], "ara"),
    ([], "ara+eng"),                  # تعذّر السؤال: الافتراضي ويقول Tesseract ما ينقص
    (["osd"], "ara+eng"),             # لا لغة نصّية مثبَّتة: لا تخمين
])
def test_ocr_uses_only_the_installed_languages(ocr, monkeypatch, installed, expected):
    monkeypatch.setattr(ocr, "languages", lambda: installed)
    assert ocr.ocr_languages() == expected


def test_ocr_languages_are_asked_once(ocr, monkeypatch):
    """السؤال عملية فرعية؛ لقطةٌ بعد لقطة لا تُعيدها."""
    calls = []
    monkeypatch.setattr(ocr, "languages", lambda: calls.append(1) or ["ara", "eng"])
    for _ in range(5):
        ocr.ocr_languages()
    assert len(calls) == 1


def test_missing_arabic_pack_is_reported_not_silent(ocr, monkeypatch, caplog):
    monkeypatch.setattr(ocr, "languages", lambda: ["eng"])
    with caplog.at_level(logging.WARNING):
        ocr.ocr_languages()
    assert any("eng" in r.getMessage() and "ناقصة" in r.getMessage()
               for r in caplog.records)


# ---------------------------------------------------------------------
# فصول الفيديو
# ---------------------------------------------------------------------
def test_chapters_report_the_sections_they_folded_away(caplog):
    from document.exporters.chapters_export import build_chapters

    plan = DocumentPlan(title="ت", sections=[
        DocumentSection(title="أول", start_timestamp=0.0),
        DocumentSection(title="قصير جدًّا", start_timestamp=4.0),
        DocumentSection(title="ثانٍ", start_timestamp=60.0),
        DocumentSection(title="ثالث", start_timestamp=120.0)])
    with caplog.at_level(logging.INFO):
        chapters = build_chapters(plan, 300.0)
    assert [t for _, t in chapters] == ["أول", "ثانٍ", "ثالث"]
    assert any("قصير جدًّا" in r.getMessage() for r in caplog.records)


def test_chapters_stay_quiet_when_nothing_was_folded(caplog):
    from document.exporters.chapters_export import build_chapters

    plan = DocumentPlan(title="ت", sections=[
        DocumentSection(title=f"قسم {i}", start_timestamp=i * 60.0) for i in range(4)])
    with caplog.at_level(logging.INFO):
        build_chapters(plan, 600.0)
    assert not any("دُمج" in r.getMessage() for r in caplog.records)


# ---------------------------------------------------------------------
# فحص بصري آلي للشرائح (LibreOffice → PDF → مواضع الكلمات)
# ---------------------------------------------------------------------
SLIDE_W_PT, SLIDE_H_PT = 13.333 * 72, 7.5 * 72
FOOTER_RULE_PT = (7.5 - 0.52) * 72
SIDE_TOLERANCE_PT = 0.5 * 72


_WORDS = ("اضغط زر الإضافة الشريط العلوي اختر الفصل المحتوى اترك الحقل فارغا النشر "
          "جميع الطلاب القائمة الجانبية المقرر التحديثات الدرجات التقارير الاختبار "
          "السؤال الإجابة الخيارات المجلد الواجب الحدث التقويم الوصف العنوان اللون "
          "الخلفية الصورة التسجيل الإعدادات الفئة الاستحقاق الإتاحة المعاينة التكرار").split()


_PREFIXES = "وبكفستلم"


def _unique_bullet(section: int, n: int) -> str:
    """نقطة **فريدة الكلمات**: تكرار النقطة نفسها يُخفي ضياع النقاط الأخيرة،
    فكلماتها موجودة في الأولى. السابقة المختلفة لكل نقطة تجعل كلماتها فريدة
    (والأرقام لا تصلح: تتبدّل مواضعها في الترتيب البصري)."""
    prefix = _PREFIXES[n % len(_PREFIXES)]
    return " ".join(prefix + _WORDS[(section * 5 + n * 7 + k) % len(_WORDS)]
                    for k in range(16))


def _stress_plan() -> DocumentPlan:
    long_bullet = _unique_bullet(0, 0)
    sections = []
    for i in range(1, 31):                       # 30 قسمًا ⇒ فهرس متعدد الشرائح
        sections.append(DocumentSection(
            title=f"القسم {i} — عنوان طويل جدًّا يتجاوز السبعين حرفًا ليختبر الالتفاف "
                  f"والاقتطاع في عنوان الشريحة وفي الفهرس معًا",
            summary="ملخّص القسم يشرح ما سيتعلّمه الطالب في هذا الجزء من الدرس بجملة طويلة نسبيًّا",
            start_timestamp=i * 60.0,
            blocks=[DocumentBlock(kind="bullets",
                                  bullets=[f"{n}. {_unique_bullet(i, n)}" for n in range(1, 8)]),
                    DocumentBlock(kind="figure", image_id=i, image_filename="wide.jpg",
                                  timestamp=i * 60.0 + 5,
                                  caption="لقطة تعرض نافذة إنشاء الحدث مع أدوات التنسيق " * 3)]
            if i % 3 else [DocumentBlock(kind="figure", image_id=i,
                                         image_filename="tall.jpg", timestamp=i * 60.0)]))
    return DocumentPlan(
        title="عنوان الدورة الطويل جدًّا الذي يمتد على أكثر من سطر في الغلاف 2026-2027",
        subtitle="دليل إجرائي مصوّر",
        abstract="ملخّص طويل " * 40,
        key_points=[long_bullet] * 6,
        sections=sections)


@pytest.fixture(scope="module")
def rendered_deck(tmp_path_factory):
    pytest.importorskip("pptx", reason="python-pptx اختيارية")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None or shutil.which("pdftotext") is None:
        pytest.skip("LibreOffice أو poppler غير متاحين")
    from PIL import Image

    from config.schemas import VideoMetadata
    from config.settings import DocumentConfig
    from document.exporters import ExportContext
    from document.exporters.pptx_export import export

    root = tmp_path_factory.mktemp("deck")
    images = root / "keyframes"
    images.mkdir()
    Image.new("RGB", (1366, 768), (40, 60, 120)).save(images / "wide.jpg")
    Image.new("RGB", (600, 1200), (120, 60, 40)).save(images / "tall.jpg")
    plan = _stress_plan()
    meta = VideoMetadata(filename="a.mp4", path=root / "a.mp4", duration_seconds=3600.0,
                         width=1920, height=1080, fps=30.0, codec="h264",
                         has_audio=True, has_video=True)
    ctx = ExportContext(plan=plan, metadata=meta, images_dir=images,
                        base_path=root / "deck.docx", document_config=DocumentConfig())
    pptx_path = export(ctx)
    profile = (root / "profile").resolve().as_uri()
    result = subprocess.run(
        [soffice, f"-env:UserInstallation={profile}", "--headless", "--convert-to",
         "pdf", "--outdir", str(root), str(pptx_path)],
        capture_output=True, text=True, timeout=600)
    pdf = root / "deck.pdf"
    if not pdf.is_file():
        pytest.skip(f"تعذّر تحويل الشرائح في هذه البيئة: {result.stderr[:120]}")
    words = subprocess.run(["pdftotext", "-bbox", str(pdf), "-"], capture_output=True,
                           text=True, check=True).stdout
    return plan, words, pptx_path


def _pages(bbox_html: str):
    import re

    for page in re.findall(r"<page [^>]*>.*?</page>", bbox_html, re.S):
        yield [(float(a), float(b), float(c), float(d), w) for a, b, c, d, w in
               re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" '
                          r'yMax="([\d.]+)">([^<]*)</word>', page)]


def test_every_slide_renders(rendered_deck):
    plan, bbox, _ = rendered_deck
    pages = list(_pages(bbox))
    figures = sum(1 for s in plan.sections for b in s.blocks if b.kind == "figure")
    agenda = -(-len(plan.sections) // 24)
    # غلاف + فهرس + ملخّص + قسم لكل + لقطة لكل + ختام
    assert len(pages) == 1 + agenda + 1 + len(plan.sections) + figures + 1


def test_no_text_leaves_the_slide(rendered_deck):
    _, bbox, _ = rendered_deck
    bad = []
    for number, words in enumerate(_pages(bbox), start=1):
        for x0, _y0, x1, y1, text in words:
            if (x0 < SIDE_TOLERANCE_PT or x1 > SLIDE_W_PT - SIDE_TOLERANCE_PT
                    or y1 > SLIDE_H_PT):
                bad.append((number, text, round(x0), round(x1), round(y1)))
    assert not bad, f"نصٌّ خارج حدود الشريحة: {bad[:5]}"


def test_body_text_never_crosses_the_footer_rule(rendered_deck):
    """النقاط الكثيرة أو الطويلة تفيض فوق التذييل — أول عطل يراه الجمهور."""
    _, bbox, _ = rendered_deck
    bad = []
    for number, words in enumerate(_pages(bbox), start=1):
        for _x0, y0, _x1, y1, text in words:
            if y0 < FOOTER_RULE_PT < y1:
                bad.append((number, text, round(y0), round(y1)))
    assert not bad, f"نصٌّ يعبر خطّ التذييل: {bad[:5]}"


def _tokens(text: str) -> set:
    import re

    import unicodedata

    # التشكيل (Mn) يُحذف: يقسم الكلمة عند ``\w`` فتختلف القطع بين الاتجاهين
    cleaned = "".join(ch for ch in text if unicodedata.category(ch) != "Mn"
                      and ch not in "\u200e\u200f\u2060")
    # بصمة الحروف لا الكلمة: ‏pdftotext يعكس العربية ويفكّ ربط «لا/لإ» بترتيبٍ
    # مختلف، فالمقارنة الحرفية تفشل على «الغلاف» و«الطلاب» وهما سليمتان.
    # الترتيب الصحيح للحروف ليس ما نفحصه — بل **وجود الكلمة** في الصفحة.
    folded = re.sub("[أإآ]", "ا", cleaned)
    return {"".join(sorted(t)) for t in re.findall(r"[\w]{3,}", folded)}


def test_no_slide_text_is_lost_off_the_slide(rendered_deck):
    """النصّ الذي يفيض عن الشريحة **يُقصّ** عند التصيير فلا يظهر في أي مكان.

    فحص «ما الذي خرج عن الحدود» لا يراه (المقصوص غير موجود في المخرج)، فالفحص
    الأصدق عكسه: كل كلمة وُضعت في الشريحة يجب أن تظهر في صفحتها المرسومة.
    الشريحة التي يُسقط مُصيِّرها بعض نقاطها عمدًا (لا تتّسع) لا تضع تلك
    النقاط في الشريحة أصلًا، فلا تُحسب مفقودة.
    """
    from pptx import Presentation

    _, bbox, pptx_path = rendered_deck
    deck = Presentation(str(pptx_path))
    pages = list(_pages(bbox))
    assert len(pages) == len(deck.slides)
    lost = []
    for number, (slide, words) in enumerate(zip(deck.slides, pages), start=1):
        wanted: set = set()
        for shape in slide.shapes:
            if shape.has_text_frame:
                wanted |= _tokens(shape.text_frame.text)
        rendered: set = set()
        for *_box, text in words:
            rendered |= _tokens(text)
        missing = wanted - rendered
        if len(missing) > max(1, len(wanted) // 50):   # تسامح 2٪ لاختلاف تقسيم الكلمات
            lost.append((number, sorted(missing)[:4]))
    assert not lost, f"شرائح فقدت نصًّا عند التصيير: {lost[:5]}"


def test_no_text_collides_with_other_text(rendered_deck):
    """نصٌّ فوق نصّ — نقطةٌ تفيض على التذييل أو عنوانٌ على ما تحته.

    فحص «العبور» لا يراه إن وقع النصّ كلّه تحت الخطّ، وفحص «الفقدان» لا يراه
    إن لم يخرج عن الصفحة؛ التصادم يراه في الحالتين.
    """
    _, bbox, _ = rendered_deck
    clashes = []
    for number, words in enumerate(_pages(bbox), start=1):
        for i, a in enumerate(words):
            for b in words[i + 1:]:
                width = min(a[2], b[2]) - max(a[0], b[0])
                height = min(a[3], b[3]) - max(a[1], b[1])
                if width <= 0 or height <= 0:
                    continue
                smaller = min((a[2] - a[0]) * (a[3] - a[1]),
                              (b[2] - b[0]) * (b[3] - b[1]))
                if smaller > 0 and width * height / smaller > 0.3:
                    clashes.append((number, a[4], b[4]))
    assert not clashes, f"نصٌّ فوق نصّ: {clashes[:4]}"
