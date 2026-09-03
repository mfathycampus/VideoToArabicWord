"""بوابات جودة المستند — بما فيها تحويل فعلي وفحص العرض، لا الـ XML وحده."""
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from docx import Document
from lxml import etree
from PIL import Image

from config.schemas import (
    AudioSegment, KeyframeMetadata, TranscriptionResult, VideoMetadata,
)
from config.settings import DocumentConfig
from document.rtl_utils import dominant_direction
from document.word_generator import DocumentGenerator
from document.planner import TimelinePlanner

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
NS = {"w": W[1:-1]}
HAS_SOFFICE = shutil.which("soffice") is not None


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    out = tmp_path_factory.mktemp("doc")
    images = out / "img"
    images.mkdir()
    keyframes = []
    for i in range(1, 4):
        name = f"{i:06d}_00-00-0{i}-000.png"
        Image.new("RGB", (1280, 720), (240 - i * 30, 240, 240)).save(images / name)
        keyframes.append(KeyframeMetadata(
            image_id=i, timestamp=float(i * 5), filename=name, scene_id=i,
            change_score=0.4, width=1280, height=720,
            selection_reason="scene_representative"))

    segments = [
        AudioSegment(id=1, start=1.0, end=4.0,
                     text_raw="النقطة الأولى في المحاضرة",
                     text_clean="النقطة الأولى في المحاضرة"),
        AudioSegment(id=2, start=6.0, end=9.0,
                     text_raw="النقطة الثانية مع الرقم 2026",
                     text_clean="النقطة الثانية مع الرقم 2026"),
        AudioSegment(id=3, start=11.0, end=14.0,
                     text_raw="الخاتمة", text_clean="الخاتمة"),
    ]
    transcript = TranscriptionResult(
        language="ar", full_text_raw="", full_text_clean="",
        segments=segments, words=[], engine={"name": "test"})
    metadata = VideoMetadata(
        filename="lecture.mp4", path=Path("lecture.mp4"),
        duration_seconds=20.0, width=1280, height=720, fps=25.0,
        codec="h264", has_audio=True, audio_sample_rate=16000,
        stored_width=1280, stored_height=720)

    matches = TimelinePlanner().build(transcript, keyframes, "lecture.mp4")
    docx = DocumentGenerator(DocumentConfig()).generate(
        matches, metadata, images, out / "doc.docx")
    return {"docx": docx, "images": images, "metadata": metadata,
            "keyframes": keyframes, "out": out}


def test_document_opens_and_keeps_all_text(built):
    doc = Document(str(built["docx"]))
    body = "\n".join(p.text for p in doc.paragraphs)
    for text in ["النقطة الأولى في المحاضرة", "النقطة الثانية مع الرقم 2026",
                 "الخاتمة"]:
        assert body.count(text) == 1


def test_all_images_embedded(built):
    """اللقطات كلها مضمّنة (بالإضافة إلى الشعار إن وُجد)."""
    from document.template import Theme
    with zipfile.ZipFile(built["docx"]) as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
    expected = len(built["keyframes"])
    assert len(media) >= expected, f"{len(media)} وسائط مقابل {expected} لقطة"


def test_images_keep_aspect_ratio(built):
    doc = Document(str(built["docx"]))
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    extents = list(doc.element.body.iter(f"{ns}ext"))
    assert extents
    source = built["images"] / built["keyframes"][0].filename
    with Image.open(source) as img:
        expected = img.height / img.width
    # الشعار مستثنى: نسبته مختلفة عن اللقطات عمدًا
    figures = [e for e in extents
               if abs(int(e.get("cy")) / int(e.get("cx")) - expected) < 0.02]
    assert len(figures) >= len(built["keyframes"]), \
        "بعض اللقطات فقدت نسبتها"


def test_document_contains_no_bidi_control_characters(built):
    """انحدار حرج: محارف العزل (U+2066/U+2069) تظهر **مربّعات في Word**.

    LibreOffice يخفيها فلا يكشف العطل؛ Word يعرضها مكتوبًا داخلها LRI و
    PDI فيمتلئ المستند بها. الاتجاه يُضبط الآن ببنية الفقرة والـ runs.
    """
    import zipfile
    with zipfile.ZipFile(built["docx"]) as archive:
        blob = b"".join(archive.read(n) for n in archive.namelist()
                        if n.endswith(".xml")).decode("utf-8", "replace")
    for char, name in (("\u2066", "LRI"), ("\u2067", "RLI"),
                       ("\u2068", "FSI"), ("\u2069", "PDI"),
                       ("\u200e", "LRM"), ("\u200f", "RLM")):
        assert char not in blob, f"محرف تحكم {name} موجود — سيظهر مربّعًا في Word"


def test_numeric_value_uses_ltr_paragraph(built):
    """القيمة الرقمية توضع في فقرة LTR فيصح ترتيبها بلا محارف مضافة."""
    import zipfile
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(built["docx"]) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    found = False
    for paragraph in root.iter(f"{W}p"):
        text = "".join(t.text or "" for t in paragraph.iter(f"{W}t"))
        if "1280×720" not in text:
            continue
        found = True
        pPr = paragraph.find(f"{W}pPr")
        bidi = pPr.find(f"{W}bidi") if pPr is not None else None
        assert bidi is not None and bidi.get(f"{W}val") == "0", \
            "قيمة رقمية في فقرة عربية — سينعكس ترتيبها"
    assert found, "قيمة الأبعاد غير موجودة"


def test_page_number_field_present(built):
    with zipfile.ZipFile(built["docx"]) as archive:
        footers = [n for n in archive.namelist() if "footer" in n and n.endswith(".xml")]
        assert footers, "لا يوجد تذييل"
        blob = b"".join(archive.read(f) for f in footers)
    assert b"PAGE" in blob, "حقل رقم الصفحة مفقود"


def test_section_is_rtl(built):
    with zipfile.ZipFile(built["docx"]) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    sectPr = next(root.iter(f"{W}sectPr"), None)
    assert sectPr is not None and sectPr.find("w:bidi", NS) is not None


def test_table_uses_bidivisual(built):
    with zipfile.ZipFile(built["docx"]) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    tblPr = next(root.iter(f"{W}tblPr"), None)
    assert tblPr is not None
    tags = [c.tag.replace(W, "") for c in tblPr]
    assert "bidiVisual" in tags and "bidi" not in tags


def test_no_dash_line_separators(built):
    """الفواصل يجب أن تكون حدود فقرات، لا سطور شرطات."""
    doc = Document(str(built["docx"]))
    for paragraph in doc.paragraphs:
        assert "-----" not in paragraph.text, "فاصل شرطات بدائي في المستند"


def test_missing_image_raises_instead_of_silent_gap(tmp_path, built):
    """المستند يجب ألا يُسلَّم صامتًا وفيه صور مفقودة."""
    from core.exceptions import DocumentGenerationError
    empty = tmp_path / "no_images"
    empty.mkdir()
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=[], words=[],
                                     engine={"name": "t"})
    matches = TimelinePlanner().build(transcript, built["keyframes"], "t")
    with pytest.raises(DocumentGenerationError):
        DocumentGenerator().generate(matches, built["metadata"], empty,
                                     tmp_path / "broken.docx")


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice غير متوفر")
def test_document_converts_to_pdf(built):
    """بوابة العرض: الـ XML الصحيح لا يكفي — الملف يجب أن يُفتح فعليًا."""
    out = built["out"]
    result = subprocess.run(
        ["soffice", "--headless", "--convert-to", "pdf",
         str(built["docx"]), "--outdir", str(out)],
        capture_output=True, timeout=240)
    pdf = out / "doc.pdf"
    assert pdf.exists() and pdf.stat().st_size > 5000, \
        f"فشل التحويل إلى PDF: {result.stderr[:300]}"


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice غير متوفر")
def test_pdf_has_expected_page_content(built):
    out = built["out"]
    pdf = out / "doc.pdf"
    if not pdf.exists():
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                        str(built["docx"]), "--outdir", str(out)],
                       capture_output=True, timeout=240)
    if not shutil.which("pdftotext"):
        pytest.skip("pdftotext غير متوفر")
    text = subprocess.run(["pdftotext", str(pdf), "-"],
                          capture_output=True, text=True, timeout=60).stdout
    assert "1280" in text and "720" in text
    # الترتيب الصحيح للأبعاد بعد العزل الاتجاهي
    assert "1280×720" in text.replace("⁦", "").replace("⁩", ""), \
        "الأبعاد معروضة معكوسة في الـ PDF"


def test_direction_detection_covers_mixed_names():
    assert dominant_direction("محاضرة 01.mp4") == "rtl"
    assert dominant_direction("lecture_01.mp4") == "ltr"
    assert dominant_direction("2026_01_15.mp4") == "ltr"
    assert dominant_direction("") == "ltr"


@pytest.mark.parametrize("size,label", [
    ((1280, 720), "أفقي"), ((360, 640), "عمودي"), ((1080, 1080), "مربع"),
    ((3840, 2160), "4K أفقي"), ((720, 1280), "عمودي هاتف"),
])
def test_image_fits_inside_page_box(tmp_path, size, label):
    """انحدار: صورة عمودية بعرض ثابت تتجاوز الصفحة فتختفي من المستند.

    الحالة الأصلية أنتجت مستندًا **بلا صور إطلاقًا** ودون أي خطأ.
    """
    from config.schemas import KeyframeMetadata, TranscriptionResult, VideoMetadata
    from docx.shared import Inches

    images = tmp_path / "img"
    images.mkdir()
    name = "000001_00-00-01-000.png"
    Image.new("RGB", size, (200, 210, 220)).save(images / name)
    keyframe = KeyframeMetadata(
        image_id=1, timestamp=1.0, filename=name, scene_id=1,
        change_score=0.5, width=size[0], height=size[1],
        selection_reason="test")
    metadata = VideoMetadata(
        filename="v.mp4", path=Path("v.mp4"), duration_seconds=10.0,
        width=size[0], height=size[1], fps=25.0, codec="h264",
        has_audio=False, stored_width=size[0], stored_height=size[1])
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=[],
                                     words=[], engine={"name": "t"})
    config = DocumentConfig()
    docx = DocumentGenerator(config).generate(
        TimelinePlanner().build(transcript, [keyframe], "t"), metadata, images,
        tmp_path / "fit.docx")

    doc = Document(str(docx))
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    extents = list(doc.element.body.iter(f"{ns}ext"))
    assert extents, f"{label}: لا توجد صورة في المستند"

    # الشعار صورة أيضًا؛ نبحث عن الصورة التي تطابق نسبة اللقطة
    expected = size[1] / size[0]
    matches = [e for e in extents
               if abs(int(e.get("cy")) / int(e.get("cx")) - expected) < 0.02]
    assert matches, f"{label}: لم تُدرج اللقطة بنسبتها الصحيحة"
    cx, cy = int(matches[0].get("cx")), int(matches[0].get("cy"))
    assert cx <= Inches(config.max_image_width_inches) + 1, \
        f"{label}: العرض يتجاوز الحد"
    assert cy <= Inches(config.max_image_height_inches) + 1, \
        f"{label}: الارتفاع يتجاوز الحد — ستُدفع الصورة خارج الصفحة"


@pytest.mark.skipif(not HAS_SOFFICE, reason="LibreOffice غير متوفر")
def test_portrait_image_actually_renders_on_page(tmp_path):
    """يتحقق أن الصورة العمودية تظهر فعلًا في الـ PDF المُصيَّر."""
    from config.schemas import KeyframeMetadata, TranscriptionResult, VideoMetadata

    images = tmp_path / "img"
    images.mkdir()
    name = "000001_00-00-01-000.png"
    Image.new("RGB", (360, 640), (40, 90, 160)).save(images / name)
    keyframe = KeyframeMetadata(image_id=1, timestamp=1.0, filename=name,
                                scene_id=1, change_score=0.5, width=360,
                                height=640, selection_reason="test")
    metadata = VideoMetadata(filename="p.mp4", path=Path("p.mp4"),
                             duration_seconds=10.0, width=360, height=640,
                             fps=25.0, codec="h264", has_audio=False,
                             stored_width=360, stored_height=640)
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=[],
                                     words=[], engine={"name": "t"})
    docx = DocumentGenerator().generate(
        TimelinePlanner().build(transcript, [keyframe], "t"), metadata, images,
        tmp_path / "portrait.docx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", str(docx),
                    "--outdir", str(tmp_path)], capture_output=True, timeout=240)
    pdf = tmp_path / "portrait.pdf"
    assert pdf.exists()
    if shutil.which("pdfimages"):
        listing = subprocess.run(["pdfimages", "-list", str(pdf)],
                                 capture_output=True, text=True, timeout=60).stdout
        assert len(listing.strip().splitlines()) > 2, \
            "الصورة العمودية غير موجودة في الـ PDF — دُفعت خارج الصفحة"


def test_mixed_direction_text_takes_arabic_path():
    """انحدار: نص فيه عربي وأغلبه لاتيني كان يخرج بلا w:rtl.

    مثال حقيقي: «تفريغ ومحتوى مرئي — static_single_slide.mp4».
    قاعدة الأغلبية صنّفته لاتينيًا فكُسر جزؤه العربي عند العرض.
    """
    from document.rtl_utils import dominant_direction
    assert dominant_direction("تفريغ ومحتوى مرئي — file_name.mp4") == "rtl"
    assert dominant_direction("محاضرة 01.mp4") == "rtl"
    assert dominant_direction("lecture_01.mp4") == "ltr"
    assert dominant_direction("1280×720") == "ltr"


def test_toc_placeholder_run_is_styled():
    """نص الحقل النائب عربي أيضًا، ويجب أن يحمل خصائص Complex Script."""
    import zipfile
    from lxml import etree
    from docx import Document as Docx
    from document.template import Theme, add_table_of_contents, build_styles

    document = Docx()
    theme = Theme()
    build_styles(document, theme)
    add_table_of_contents(document, theme)

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    ns = {"w": W[1:-1]}
    for run in document.element.body.iter(f"{W}r"):
        text = "".join(t.text or "" for t in run.iter(f"{W}t"))
        if not any("؀" <= ch <= "ۿ" for ch in text):
            continue
        rPr = run.find("w:rPr", ns)
        assert rPr is not None and rPr.find("w:rtl", ns) is not None, \
            f"حقل بلا اتجاه: {text[:40]}"


def test_document_has_navigable_headings():
    """الفهرس وجزء التنقّل يعملان بالأنماط لا بالتنسيق اليدوي."""
    from config.schemas import (AudioSegment, TranscriptionResult, VideoMetadata)
    from document.planner import TimelinePlanner
    import tempfile

    segments = [AudioSegment(id=i, start=i * 20.0, end=i * 20.0 + 15.0,
                             text_raw=f"فقرة رقم {i} من الشرح.",
                             text_clean=f"فقرة رقم {i} من الشرح.")
                for i in range(1, 13)]
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=segments,
                                     words=[], engine={"name": "t"})
    metadata = VideoMetadata(filename="v.mp4", path=Path("v.mp4"),
                             duration_seconds=240.0, width=1280, height=720,
                             fps=25.0, codec="h264", has_audio=True,
                             stored_width=1280, stored_height=720)
    plan = TimelinePlanner().build(transcript, [], "محاضرة تجريبية")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "nav.docx"
        DocumentGenerator().generate(plan, metadata, Path(tmp), out)
        doc = Document(str(out))
        headings = [p for p in doc.paragraphs
                    if p.style.name.startswith("Heading")]
        assert len(headings) >= 2, "لا عناوين — الفهرس لن يعمل"


def test_toc_field_is_present():
    import zipfile
    from config.schemas import (AudioSegment, TranscriptionResult, VideoMetadata)
    from document.planner import TimelinePlanner
    import tempfile

    segments = [AudioSegment(id=1, start=0.0, end=5.0, text_raw="نص.",
                             text_clean="نص.")]
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=segments,
                                     words=[], engine={"name": "t"})
    metadata = VideoMetadata(filename="v.mp4", path=Path("v.mp4"),
                             duration_seconds=10.0, width=1280, height=720,
                             fps=25.0, codec="h264", has_audio=True,
                             stored_width=1280, stored_height=720)
    plan = TimelinePlanner().build(transcript, [], "عنوان")
    with tempfile.TemporaryDirectory() as tmp:
        out = Path(tmp) / "toc.docx"
        DocumentGenerator().generate(plan, metadata, Path(tmp), out)
        with zipfile.ZipFile(out) as archive:
            body = archive.read("word/document.xml")
            # حقول الترقيم تعيش في جزء التذييل لا في المتن
            footers = b"".join(archive.read(n) for n in archive.namelist()
                               if "footer" in n and n.endswith(".xml"))
    assert b"TOC" in body, "حقل الفهرس مفقود"
    assert b"NUMPAGES" in footers, "حقل عدد الصفحات مفقود من التذييل"
    assert b"PAGE" in footers


def test_logo_is_embedded_when_available(built):
    """الشعار يظهر في الغلاف والرأس متى وُجد في assets/."""
    import zipfile
    from document.template import Theme
    if Theme().resolved_logo() is None:
        pytest.skip("لا يوجد شعار في المشروع")
    with zipfile.ZipFile(built["docx"]) as archive:
        names = archive.namelist()
        media = [n for n in names if n.startswith("word/media/")]
        headers = [n for n in names if "header" in n and n.endswith(".xml")]
    assert media, "لا وسائط في المستند"
    assert headers, "لا رأس صفحة"


def test_missing_logo_does_not_break_generation(tmp_path):
    """غياب الشعار يجب ألا يُفشل التوليد."""
    from config.schemas import (AudioSegment, TranscriptionResult, VideoMetadata)
    from document.planner import TimelinePlanner
    from document.template import Theme

    theme = Theme(logo_path=tmp_path / "nope.png")
    assert theme.resolved_logo() is None or theme.resolved_logo().exists()

    segments = [AudioSegment(id=1, start=0.0, end=5.0, text_raw="نص.",
                             text_clean="نص.")]
    transcript = TranscriptionResult(language="ar", full_text_raw="",
                                     full_text_clean="", segments=segments,
                                     words=[], engine={"name": "t"})
    metadata = VideoMetadata(filename="v.mp4", path=Path("v.mp4"),
                             duration_seconds=10.0, width=1280, height=720,
                             fps=25.0, codec="h264", has_audio=True,
                             stored_width=1280, stored_height=720)
    plan = TimelinePlanner().build(transcript, [], "عنوان")
    out = tmp_path / "nologo.docx"
    DocumentGenerator(theme=theme).generate(plan, metadata, tmp_path, out)
    assert out.exists()


def test_arabic_date_avoids_reversible_numeric_groups():
    """انحدار: «2026-08-31» يُعرض «31-08-2026» — تاريخ مختلف للقارئ."""
    from datetime import datetime
    from utils.timestamps import arabic_datetime
    text = arabic_datetime(datetime(2026, 8, 31, 21, 33))
    assert "أغسطس" in text and "2026" in text and "31" in text
    assert "2026-08-31" not in text, "صيغة ISO تنعكس بصريًا في العربية"


def test_ocr_text_renders_under_the_figure_when_present(tmp_path):
    """نص الشاشة (OCR) على لقطة يجب أن يظهر في المستند تحتها، منفصلًا
    عن التسمية العادية — انظر document/word_generator.py::_add_figure."""
    images = tmp_path / "img"
    images.mkdir()
    name = "000001_00-00-01-000.png"
    Image.new("RGB", (1280, 720), (230, 230, 230)).save(images / name)
    keyframe = KeyframeMetadata(
        image_id=1, timestamp=1.0, filename=name, scene_id=1,
        change_score=0.4, width=1280, height=720,
        selection_reason="scene_representative",
        ocr_text="تعريف: المشتقة هي معدل التغيّر اللحظي")

    segments = [AudioSegment(id=1, start=0.0, end=2.0,
                             text_raw="كلام", text_clean="كلام")]
    transcript = TranscriptionResult(
        language="ar", full_text_raw="", full_text_clean="",
        segments=segments, words=[], engine={"name": "test"})
    metadata = VideoMetadata(
        filename="lecture.mp4", path=Path("lecture.mp4"),
        duration_seconds=5.0, width=1280, height=720, fps=25.0,
        codec="h264", has_audio=True, audio_sample_rate=16000,
        stored_width=1280, stored_height=720)

    plan = TimelinePlanner().build(transcript, [keyframe], "lecture.mp4")
    docx = DocumentGenerator(DocumentConfig()).generate(
        plan, metadata, images, tmp_path / "ocr_doc.docx")

    doc = Document(str(docx))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "النص الظاهر على الشاشة" in body
    assert "تعريف: المشتقة هي معدل التغيّر اللحظي" in body


def test_absent_ocr_text_adds_no_extra_paragraph(tmp_path):
    """لا نص OCR ⇐ لا فقرة إضافية ولا تسمية «النص الظاهر على الشاشة» —
    الميزة صامتة تمامًا حين لا بيانات لديها."""
    images = tmp_path / "img"
    images.mkdir()
    name = "000001_00-00-01-000.png"
    Image.new("RGB", (1280, 720), (230, 230, 230)).save(images / name)
    keyframe = KeyframeMetadata(
        image_id=1, timestamp=1.0, filename=name, scene_id=1,
        change_score=0.4, width=1280, height=720,
        selection_reason="scene_representative")

    segments = [AudioSegment(id=1, start=0.0, end=2.0,
                             text_raw="كلام", text_clean="كلام")]
    transcript = TranscriptionResult(
        language="ar", full_text_raw="", full_text_clean="",
        segments=segments, words=[], engine={"name": "test"})
    metadata = VideoMetadata(
        filename="lecture.mp4", path=Path("lecture.mp4"),
        duration_seconds=5.0, width=1280, height=720, fps=25.0,
        codec="h264", has_audio=True, audio_sample_rate=16000,
        stored_width=1280, stored_height=720)

    plan = TimelinePlanner().build(transcript, [keyframe], "lecture.mp4")
    docx = DocumentGenerator(DocumentConfig()).generate(
        plan, metadata, images, tmp_path / "no_ocr_doc.docx")

    doc = Document(str(docx))
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "النص الظاهر على الشاشة" not in body
