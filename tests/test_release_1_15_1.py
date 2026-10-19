"""1.15.1 — ما رُصد على «دليل إعداد وإدارة الاختبارات الإلكترونية».

دليلٌ مصوّر حقيقي من تسجيل 40 دقيقة، ولوحة إحصاءات على الجهاز نفسه:

* فهرس الأشكال 36 سطرًا من «شكل N · التوقيت 00:0x:xx» بلا وصف.
* قسمان متتاليان يشرحان الخطوات والملاحظة نفسيهما.
* شكلان بلا تعليق، وتذييلٌ يسمّي الدليل «تقرير تفريغ الفيديو»، ومقاس Letter.
* أسماء طلاب ظاهرة («meshal, Alawaad»، و«aljamal» بجانب اسمٍ مموَّه).
* شريط مهام ويندوز في كل لقطة.
* آخر شكل عند 17:18 من 40:00 ولا تحذير.
* الإحصاءات: الجارية «لم تكتمل»، و«5 من 7» بجانب «83٪»، و«فشلت · جودة 72»،
  و«4 مستند»، ومحور «1 2 3 4 5»، و«سرعة المعالجة —».
* الشعار: لا اختيار إلا مسارٌ يضيع، ولا يسري على الويب والشرائح.
"""
from __future__ import annotations

import json
import zipfile
from datetime import datetime
from pathlib import Path

import numpy as np
import pytest
from PIL import Image

from config.schemas import (
    AudioSegment,
    DocumentBlock,
    DocumentPlan,
    DocumentSection,
    VideoMetadata,
)
from config.settings import AppConfig, DocumentConfig


# ── المستند ─────────────────────────────────────────────────────────
def _guide_plan(figures: int = 22) -> DocumentPlan:
    blocks = []
    for i in range(1, figures + 1):
        blocks.append(DocumentBlock(kind="step", text=f"اضغط الزرّ رقم {i}."))
        blocks.append(DocumentBlock(kind="figure", image_id=i, image_filename="f.jpg",
                                    timestamp=10.0 * i,
                                    caption=f"نافذة نسخ الاختبار رقم {i}"))
    plan = DocumentPlan(title="دليل إعداد الاختبارات", subtitle="دليل إجرائي مصوّر",
                        sections=[DocumentSection(title=f"قسم {k}", level=1,
                                                  blocks=blocks if k == 1 else
                                                  [DocumentBlock(kind="step", text="خطوة")])
                                  for k in range(1, 4)])
    plan.generated_by = "ai:fake/model+vision"
    plan.total_figures_out = figures
    return plan


def _docx(tmp_path: Path, plan: DocumentPlan, config: DocumentConfig | None = None) -> Path:
    from document.word_generator import DocumentGenerator

    images = tmp_path / "keyframes"
    images.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (1366, 683), (200, 210, 220)).save(images / "f.jpg")
    metadata = VideoMetadata(filename="اعداد الاختبارات.mp4", path=tmp_path / "x.mp4",
                             duration_seconds=2400.0, width=1366, height=683, fps=25.0,
                             codec="h264", has_audio=True, has_video=True)
    out = tmp_path / "guide.docx"
    import document.word_generator as wg
    original = wg.assert_lossless
    wg.assert_lossless = lambda plan: None
    try:
        DocumentGenerator(config or DocumentConfig()).generate(plan, metadata, images, out)
    finally:
        wg.assert_lossless = original
    return out


def _xml(path: Path, part: str = "word/document.xml") -> str:
    with zipfile.ZipFile(path) as archive:
        return archive.read(part).decode("utf-8")


def test_caption_paragraph_that_feeds_the_figure_index_carries_the_description(tmp_path):
    """فهرس الأشكال يُبنى من فقرة حقل SEQ — فيجب أن تحمل الوصف لا التوقيت."""
    from docx import Document

    document = Document(str(_docx(tmp_path, _guide_plan())))
    seq = [p for p in document.paragraphs
           if "SEQ Figure" in p._p.xml and "TOC" not in p._p.xml]
    assert len(seq) == 22
    for paragraph in seq:
        assert "نافذة نسخ الاختبار" in paragraph.text
        assert "التوقيت" not in paragraph.text
    stamps = [p for p in document.paragraphs if p.text.startswith("التوقيت ")]
    assert len(stamps) == 22                  # التوقيت باقٍ، في سطرٍ خارج الفهرس


def test_footer_names_the_document_type_and_page_is_a4(tmp_path):
    out = _docx(tmp_path, _guide_plan(3))
    footers = "".join(_xml(out, name) for name in zipfile.ZipFile(out).namelist()
                      if name.startswith("word/footer"))
    assert "دليل إجرائي مصوّر" in footers
    assert "تقرير تفريغ الفيديو" not in footers
    body = _xml(out)
    assert 'w:w="11906"' in body and 'w:h="16838"' in body      # A4 بالـ twips


def test_custom_footer_text_is_kept(tmp_path):
    out = _docx(tmp_path, _guide_plan(3), DocumentConfig(footer_text="إدارة التقنية"))
    footers = "".join(_xml(out, name) for name in zipfile.ZipFile(out).namelist()
                      if name.startswith("word/footer"))
    assert "إدارة التقنية" in footers


# ── المؤلّف المرئي ───────────────────────────────────────────────────
def test_adjacent_sections_repeating_one_task_are_merged():
    from ai.visual_author import merge_repeated_sections

    first = DocumentSection(title="تحديد تاريخ ووقت الاستحقاق لكل فصل", level=1, blocks=[
        DocumentBlock(kind="step", text="افتح حقل «تاريخ الاستحقاق» واختر اليوم من التقويم."),
        DocumentBlock(kind="note", text="ضبط الوقت بالساعة مهم لأن الاحتساب يتم بالساعة وليس باليوم.")])
    second = DocumentSection(title="نسخ الاختبار إلى فصل آخر وتحديد موعد الاستحقاق", level=1, blocks=[
        DocumentBlock(kind="step", text="اختر اليوم والشهر والسنة لموعد الاختبار في الفصل الجديد."),
        DocumentBlock(kind="note", text="ضبط الوقت بالساعة مهم لأن الاحتساب يتم بالساعة وليس باليوم فقط.")])
    merged = merge_repeated_sections([first, second])
    assert len(merged) == 1
    notes = [b for b in merged[0].blocks if b.kind == "note"]
    assert len(notes) == 1                    # الملاحظة المكرّرة لا تُعاد


def test_neighbouring_distinct_tasks_are_not_merged():
    from ai.visual_author import merge_repeated_sections

    titles = ["إنشاء الاختبار وتسميته", "ضبط إعدادات الاختبار العامة",
              "ضبط شريط أدوات الاختبار وإعدادات الطلاب", "التعرف على أنواع الأسئلة",
              "نسخ الاختبار إلى فصل آخر وتحديد موعد الاستحقاق",
              "التحقق من نجاح نسخ الاختبار", "متابعة نتائج الطلاب في الاختبار",
              "تحديد نقاط أسئلة الاختبار"]
    sections = [DocumentSection(title=t, level=1, blocks=[
        DocumentBlock(kind="step", text=f"اضغط {t[::-1]}")]) for t in titles]
    assert len(merge_repeated_sections(sections)) == len(titles)


def test_figure_restored_without_caption_gets_one_from_its_step():
    from ai.visual_author import _fallback_caption

    assert _fallback_caption("3. اضغط «إضافة محتوى» ثم اختر «اختبار».", "قسم") == \
        "اضغط «إضافة محتوى» ثم اختر «اختبار»"
    assert _fallback_caption("", "إنشاء الاختبار") == "إنشاء الاختبار"


def test_later_parts_see_the_sections_already_written(tmp_path):
    from tests.test_visual_author import FakeVision, _frames, _transcript

    from ai.visual_author import VisualAuthor, VisualAuthorConfig

    frames = _frames(tmp_path, 4)

    def part(title, ids):
        return {"sections": [{"title": title, "blocks": [
            {"type": "figure", "id": i, "caption": f"لقطة {i}"} for i in ids]}]}
    fake = FakeVision([part("تحديد تاريخ الاستحقاق", [1, 2]),
                       part("التحقق من النسخ", [3, 4]),
                       {"title": "دليل", "abstract": "", "key_points": []}])
    VisualAuthor(fake, VisualAuthorConfig(max_images_per_call=2)).build_plan(
        _transcript([(5.0, "بداية")]), frames, tmp_path, "ملف")
    second_intro = fake.calls[1][1][0]["text"]
    assert "تحديد تاريخ الاستحقاق" in second_intro
    assert "تحديد تاريخ الاستحقاق" not in fake.calls[0][1][0]["text"]


# ── إخفاء الأسماء ────────────────────────────────────────────────────
def _w(text, line, left, top=0, height=14, width=None):
    return {"text": text, "conf": 90.0, "line": ("1", "1", str(line)),
            "left": left, "top": top, "width": width or 8 * len(text), "height": height}


def test_fifth_word_of_a_long_name_is_hidden_too():
    from video.redact import find_name_boxes

    words = [_w("Noura", 1, 600), _w("Saleh", 1, 650), _w("Mohammed", 1, 700),
             _w("Ali", 1, 780), _w("aljamal", 1, 815)]
    boxes, _ = find_name_boxes(words)
    assert {b["text"] for b in boxes} == {"Noura", "Saleh", "Mohammed", "Ali", "aljamal"}


def test_right_aligned_name_rows_in_an_rtl_table_are_hidden():
    """جدول نتائج عربي: الأسماء تنتهي عند الحافة نفسها وتبدأ من مواضع مختلفة."""
    from video.redact import find_name_boxes

    # الصفّ الأول يُكشف اسمًا، والثاني قراءةٌ مشوّهة تنتهي عند الحافة نفسها
    words = [_w("Noura", 1, 800, top=340, width=50), _w("Alawaad", 1, 856, top=340, width=72),
             _w("rneshal,", 2, 860, top=430, width=60), _w("Alawaxd", 2, 922, top=430, width=6)]
    boxes, _ = find_name_boxes(words)
    assert {"rneshal,", "Alawaxd"} <= {b["text"] for b in boxes}


def test_name_truncated_at_its_start_is_hidden():
    """شريط المنصّة يقصّ الاسم من أوّله: «...d Yosef»."""
    from video.redact import looks_like_person_name

    assert looks_like_person_name("...d Yosef")
    assert looks_like_person_name("…mmed")


# ── قصّ شريط المهام ──────────────────────────────────────────────────
def _screen(seed: int, height=683, width=1366, taskbar=35):
    rng = np.random.default_rng(seed)
    # محتوى يتغيّر صفًّا صفًّا بين الإطارات (صفحاتٌ تُمرَّر وتُفتح)
    rows = rng.integers(0, 255, (height, 1, 1), dtype=np.uint8)
    frame = np.repeat(np.repeat(rows, width, axis=1), 3, axis=2)
    frame[height - taskbar:] = 32                       # شريط المهام ثابت
    frame[height - taskbar:height - taskbar + 6] += np.uint8(seed % 5)   # حافّة Mica
    return frame


def test_taskbar_is_cropped_even_when_the_recording_ends_static():
    from video.screen_crop import detect

    moving = [_screen(i) for i in range(4)]
    frozen = [moving[-1].copy() for _ in range(6)]      # التسجيل يدور بعد انتهاء الشرح
    box = detect(moving + frozen)
    assert box.bottom >= 33


def test_mostly_static_frames_without_a_taskbar_edge_are_not_cropped():
    from video.screen_crop import detect

    base = np.full((683, 1366, 3), 240, dtype=np.uint8)
    frames = []
    for i in range(4):
        frame = base.copy()
        frame[300:320, 100:200] = 20 * i                 # تغيّرٌ صغير وسط الصفحة
        frames.append(frame)
    box = detect(frames)
    assert box.bottom == 0 and box.top == 0


# ── مدى المستند من التسجيل ───────────────────────────────────────────
def test_document_ending_long_before_the_recording_is_flagged():
    from document.quality import content_span

    plan = _guide_plan(3)
    for block in plan.sections[0].blocks:
        if block.image_id:
            block.timestamp = 1038.0                     # 17:18
    segments = [AudioSegment(id=1, start=1100.0, end=1900.0, text_raw="كلام", text_clean="كلام")]
    info, warnings = content_span(plan, [], segments, 2400.0)
    assert info["last_figure_seconds"] == 1038.0
    assert warnings and "17:18" in warnings[0] and "40:00" in warnings[0]


def test_document_covering_the_recording_is_not_flagged():
    from document.quality import content_span

    plan = _guide_plan(3)
    plan.sections[0].blocks[-1].timestamp = 2300.0
    assert content_span(plan, [], [], 2400.0)[1] == []


# ── الإحصاءات ────────────────────────────────────────────────────────
def _fact(status, at, quality=80.0, outputs=(".docx",), minutes=40):
    from core.stats import JobFact

    return JobFact(key=at, at=at, status=status, name="م", duration_seconds=minutes * 60,
                   quality_score=quality, outputs=list(outputs))


def test_running_job_is_not_counted_as_unfinished_and_denominators_agree():
    from core.stats import compute

    facts = [_fact("completed", f"2026-10-0{i}T10:00:00") for i in range(1, 6)]
    facts += [_fact("failed", "2026-10-15T10:00:00", quality=72.0),
              _fact("running", "2026-10-18T08:00:00", quality=None)]
    busy = compute(facts, now=datetime(2026, 10, 18), busy=True)
    assert busy.running == 1 and busy.unfinished == 0
    assert busy.finished == 6 and round(busy.success_rate) == 83
    idle = compute(facts, now=datetime(2026, 10, 18), busy=False)
    assert idle.running == 0 and idle.unfinished == 1
    assert busy.quality_dates[-1].startswith("2026-10-05")
    assert len(busy.quality_dates) == len(busy.quality_trend) == 5


def test_speed_is_inferred_from_job_state_for_jobs_before_the_ledger(tmp_path):
    from core.stats import _inferred_elapsed

    job = tmp_path / "job"
    job.mkdir()
    (job / "job_state.json").write_text(json.dumps({
        "created_at": "2026-10-05T10:00:00+00:00",
        "updated_at": "2026-10-05T10:51:00+00:00"}), encoding="utf-8")
    assert _inferred_elapsed(job) == 51 * 60
    (job / "job_state.json").write_text(json.dumps({
        "created_at": "2026-10-01T10:00:00+00:00",
        "updated_at": "2026-10-05T10:51:00+00:00"}), encoding="utf-8")
    assert _inferred_elapsed(job) == 0.0                 # استُؤنفت بعد أيام


@pytest.mark.parametrize("n, expected", [
    (1, "مستند واحد"), (2, "مستندان"), (4, "4 مستندات"), (10, "10 مستندات"),
    (11, "11 مستندًا"), (103, "103 مستندات"), (2.5, "2.5 مستند")])
def test_arabic_counted_nouns(n, expected):
    from utils.arabic_plural import DOCUMENTS, count_label

    assert count_label(n, DOCUMENTS) == expected


def test_minutes_keep_one_decimal_below_ten():
    from utils.arabic_plural import minutes_label

    assert minutes_label(160) == "2.7 دقيقة"
    assert minutes_label(89 * 60) == "89 دقيقة"
    assert minutes_label(3 * 60) == "3 دقائق"


# ── الشعار ───────────────────────────────────────────────────────────
def test_logo_choices_resolve(tmp_path):
    from document.branding import resolve_logo

    assert resolve_logo(DocumentConfig()).name == "logo.png"
    assert resolve_logo(DocumentConfig(logo_choice="wordmark")).name == "logo_wordmark.png"
    assert resolve_logo(DocumentConfig(logo_choice="none")) is None
    missing = DocumentConfig(logo_choice="custom", logo_path=tmp_path / "gone.png")
    assert resolve_logo(missing).name == "logo.png"      # مفقود ← شعار معين، لا انهيار


def test_legacy_logo_path_becomes_custom_choice(tmp_path):
    logo = tmp_path / "school.png"
    Image.new("RGB", (300, 120), (10, 80, 160)).save(logo)
    config = DocumentConfig(logo_path=logo)
    assert config.logo_choice == "custom"
    assert DocumentConfig(logo_choice="bogus").logo_choice == "app"


def test_custom_logo_is_copied_so_moving_the_original_does_not_lose_it(tmp_path):
    from document.branding import import_custom_logo

    source = tmp_path / "Downloads" / "school.jpg"
    source.parent.mkdir()
    Image.new("RGB", (2400, 800), (10, 80, 160)).save(source)
    copy = import_custom_logo(source, tmp_path / "branding")
    source.unlink()
    with Image.open(copy) as image:
        assert image.format == "PNG" and max(image.size) <= 1200
    with pytest.raises(ValueError):
        bad = tmp_path / "not-image.png"
        bad.write_text("x")
        import_custom_logo(bad, tmp_path / "branding")


def test_no_logo_choice_removes_it_from_word(tmp_path):
    with_logo = _xml(_docx(tmp_path / "a", _guide_plan(3)))
    without = _xml(_docx(tmp_path / "b", _guide_plan(3), DocumentConfig(logo_choice="none")))
    assert with_logo.count("<pic:pic") == without.count("<pic:pic") + 1


def test_logo_reaches_the_web_page_and_respects_none(tmp_path):
    from tests.test_exporters import _metadata, _plan

    from document.exporters import ExportContext
    from document.exporters.html_export import render_html

    def page(config):
        images = tmp_path / "keyframes"
        images.mkdir(exist_ok=True)
        return render_html(ExportContext(plan=_plan(), metadata=_metadata(tmp_path),
                                         images_dir=images, base_path=tmp_path / "l.docx",
                                         document_config=config))
    assert '<img class="brand"' in page(DocumentConfig())
    assert '<img class="brand"' not in page(DocumentConfig(logo_choice="none"))
    assert "#2D2E82" not in page(DocumentConfig())        # لون المدرسة القديم


def test_logo_choice_survives_save_and_load(tmp_path):
    config = AppConfig()
    config.document.logo_choice = "wordmark"
    path = tmp_path / "config.yaml"
    config.save(path)
    assert AppConfig.load(path).document.logo_choice == "wordmark"


def test_lone_lowercase_surname_in_the_names_column_is_hidden():
    """صفّ «aljamal» وحده تحت «meshal, Alawaad» في عمود الأسماء نفسه."""
    from video.redact import find_name_boxes

    words = [_w("aljamal", 1, 937, top=340, width=56),
             _w("meshal,", 2, 864, top=428, width=60), _w("Alawaad", 2, 930, top=428, width=63)]
    boxes, _ = find_name_boxes(words)
    assert {"aljamal", "meshal,", "Alawaad"} <= {b["text"] for b in boxes}
